import json
import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document
from pydantic import ValidationError
from pypdf import PdfReader

from resume_app.llm_workflow import APPLICATION_FIELDS, PROMPT_VERSION, apply_plan
from resume_app.models import Publication, ResumeProfile, TailoringPlan, resume_cache_data
from resume_app.pdf import template_docx, template_pdf
from resume_app.pdf.storage import store_pdf
from resume_app.storage import fingerprint
from test_product import JOB, PLAN, SOURCE, ProductTestCase, word_bytes


PAPERS = [
    {
        "title": "Grounded Resume Systems",
        "authors": ["J. Example", "张三"],
        "venue": "Sample Conference", "date": "2025",
        "doi": "10.1234/resume.1", "url": "https://example.org/papers/one",
        "details": ["Volume 12, pages 1-9", "Equal contribution noted in the source"],
    },
    {
        "title": "Earlier Research", "authors": ["A. First"], "venue": "Sample Journal",
        "date": "2023", "doi": "", "url": "",
        "details": ["Preprint identifier: arXiv:2301.00001"],
    },
]


def research_profile():
    return ResumeProfile.model_validate({**SOURCE.model_dump(), "publications": PAPERS})


def word_text(content):
    document = Document(BytesIO(content))
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [paragraph.text for table in document.tables for row in table.rows
           for cell in row.cells for paragraph in cell.paragraphs]
    )


class PublicationSchemaTests(unittest.TestCase):
    def test_old_profiles_gain_an_empty_section_without_changing_the_cache_shape(self):
        legacy = SOURCE.model_dump(exclude={"publications"})
        normalized = ResumeProfile.model_validate(legacy)
        self.assertEqual(normalized.publications, [])
        self.assertEqual(resume_cache_data(normalized), legacy)

    def test_partial_citations_preserve_raw_source_without_inventing_metadata(self):
        raw = "A. Author et al. Unsplit research citation, preprint, 2024."
        paper = Publication(details=[raw])
        self.assertEqual(paper.title, "")
        self.assertEqual(paper.authors, [])
        self.assertEqual(paper.doi, "")
        self.assertEqual(paper.details, [raw])

    def test_empty_or_unrecognized_publication_fields_are_rejected(self):
        for paper in ({}, {"authors": [""]}, {"details": [" "]}, {"title": "Paper", "invented_status": "published"}):
            with self.subTest(paper=paper), self.assertRaises(ValidationError):
                Publication.model_validate(paper)

    def test_publication_limit_is_one_hundred_and_does_not_truncate(self):
        self.assertEqual(len(ResumeProfile.model_validate({"publications": [PAPERS[0]] * 100}).publications), 100)
        with self.assertRaises(ValidationError):
            ResumeProfile.model_validate({"publications": [PAPERS[0]] * 101})

    def test_tailoring_preserves_every_publication_fact_and_order(self):
        source = research_profile()
        result = apply_plan(source, TailoringPlan.model_validate(PLAN))
        self.assertEqual(result.publications, source.publications)
        result.publications[0].authors.append("Manual edit")
        self.assertEqual(source.publications[0].authors, ["J. Example", "张三"])
        with self.assertRaises(ValidationError):
            TailoringPlan.model_validate({**PLAN, "publications": []})


class PublicationWorkflowTests(ProductTestCase):
    def setUp(self):
        super().setUp()
        self.setup_user()
        self.settings()

    def test_uploaded_document_prompt_and_schema_explicitly_cover_papers(self):
        text = "Sample Applicant\nPublications\nJ. Example, 张三. Grounded Resume Systems. Sample Conference, 2025. DOI: 10.1234/resume.1"
        with patch.object(self.workflow, "_complete", return_value=research_profile().model_dump_json()) as model:
            response = self.post("/api/resume/upload", data={
                "file": (BytesIO(word_bytes(text)), "research.docx"), "consent": "yes",
            })
            self.assertEqual(response.status_code, 202)
            self.assertEqual(self.wait_task(response.get_json()["task_id"])["status"], "succeeded")
        payload = json.loads(model.call_args.args[1])
        self.assertIn("Grounded Resume Systems", payload["resume_text"])
        self.assertIn("publications", payload["schema"]["properties"])
        self.assertIn("Publication", payload["schema"]["$defs"])
        self.assertIn("authorship order", model.call_args.args[2])
        self.assertIn("do not look them up", model.call_args.args[2])
        record = self.client.get("/api/state").get_json()["resume"]
        self.assertEqual(record["pending_import"]["profile"]["publications"], PAPERS)
        self.assertEqual(self.post("/api/resume/import/apply", {
            "edit_token": record["edit_token"], "import_id": record["pending_import"]["id"],
            "profile": record["pending_import"]["profile"],
        }).status_code, 200)
        record = self.client.get("/api/state").get_json()["resume"]
        self.assertEqual(record["draft"]["publications"], PAPERS)
        self.assertEqual(self.post("/api/resume/confirm", {
            "edit_token": record["edit_token"], "profile": record["draft"],
        }).status_code, 200)
        self.assertEqual(self.client.get("/api/state").get_json()["resume"]["confirmed"]["publications"], PAPERS)

    def test_legacy_api_profiles_are_normalized_without_rewriting_stored_data(self):
        legacy = SOURCE.model_dump(exclude={"publications"})
        self.database.resumes.insert_one({
            "_id": "local", "confirmed": legacy, "edit_token": "legacy-token",
            "confirmed_hash": fingerprint(legacy), "draft": None,
        })
        record = self.client.get("/api/state").get_json()["resume"]
        self.assertEqual(record["confirmed"]["publications"], [])
        self.assertNotIn("publications", self.database.resumes.find_one({"_id": "local"})["confirmed"])
        self.assertEqual(self.post("/api/resume/confirm", {
            "edit_token": record["edit_token"], "profile": record["confirmed"],
        }).status_code, 200)
        self.assertEqual(self.database.resumes.find_one({"_id": "local"})["confirmed_hash"], fingerprint(legacy))

    def test_no_papers_does_not_trigger_a_paid_regeneration_of_a_legacy_cache(self):
        settings = self.workflow.llm_settings("local")
        legacy = SOURCE.model_dump(exclude={"publications"})
        old_id = fingerprint({
            "user_id": "local", "job": JOB.model_dump(exclude=APPLICATION_FIELDS),
            "source": legacy, "llm": settings.model_dump(exclude={"api_key"}),
            "prompt_version": PROMPT_VERSION,
        })
        self.database.tailored_resumes.insert_one({"_id": old_id, "user_id": "local"})
        with patch.object(self.workflow, "_complete") as model:
            resume_id, cached = self.workflow.tailor("local", JOB, SOURCE, settings)
        self.assertTrue(cached)
        self.assertEqual(resume_id, old_id)
        model.assert_not_called()

    def test_adding_papers_changes_cache_and_preserves_them_in_the_job_resume(self):
        settings = self.workflow.llm_settings("local")
        with patch.object(self.workflow, "_complete", return_value=json.dumps(PLAN)) as model:
            old_id, _ = self.workflow.tailor("local", JOB, SOURCE, settings)
            new_id, cached = self.workflow.tailor("local", JOB, research_profile(), settings)
        self.assertFalse(cached)
        self.assertNotEqual(new_id, old_id)
        self.assertEqual(model.call_count, 2)
        self.assertIn("Publications are factual records", model.call_args.args[2])
        stored = self.database.tailored_resumes.find_one({"_id": new_id})
        self.assertEqual(stored["resume"]["publications"], PAPERS)
        self.assertEqual(stored["source_resume"]["publications"], PAPERS)

    def test_publications_only_resume_can_be_confirmed(self):
        profile = ResumeProfile.model_validate({"contact": {"name": "Researcher"}, "publications": PAPERS})
        self.database.resumes.insert_one({"_id": "local", "draft": profile.model_dump(), "edit_token": "papers-only"})
        self.assertEqual(self.post("/api/resume/confirm", {
            "edit_token": "papers-only", "profile": profile.model_dump(),
        }).status_code, 200)

    def test_normalizing_a_legacy_job_response_keeps_its_confirmed_pdf_valid(self):
        legacy = SOURCE.model_dump(exclude={"publications"})
        content, template = template_pdf(ResumeProfile.model_validate(legacy), "simple")
        artifact = store_pdf(
            self.database, "local", "legacy-paper-pdf", "legacy-file",
            fingerprint(legacy), content, template=template,
        )
        job_id = self.workflow.save_job("local", JOB)
        self.database.tailored_resumes.insert_one({
            "_id": "legacy-paper-pdf", "user_id": "local", "job_id": job_id,
            "job": JOB.model_dump(), "resume": legacy, "source_resume": legacy,
            "review_status": "reviewed", "edit_token": "legacy-token",
            "pdf": artifact, "template_id": "simple",
        })
        result = self.client.get("/api/tailored/legacy-paper-pdf").get_json()
        self.assertEqual(result["resume"]["publications"], [])
        self.assertEqual(result["source_resume"]["publications"], [])
        self.assertTrue(result["pdf_ready"])
        self.assertEqual(self.client.get("/api/tailored/legacy-paper-pdf/pdf").data, content)
        self.assertNotIn("publications", self.database.tailored_resumes.find_one({"_id": "legacy-paper-pdf"})["resume"])


class PublicationExportTests(unittest.TestCase):
    def test_both_pdf_and_word_templates_include_complete_publication_details(self):
        profile = research_profile()
        for template in ("simple", "classic-044"):
            with self.subTest(template=template):
                content, _ = template_pdf(profile, template)
                text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
                for value in ("Publications" if template == "simple" else "PUBLICATIONS",
                              "Grounded Resume Systems", "Earlier Research", "J. Example", "张三",
                              "Sample Conference", "2025", "10.1234/resume.1",
                              "https://example.org/papers/one", "Volume 12, pages 1-9"):
                    self.assertIn(value, text)
                self.assertLess(text.index("Grounded Resume Systems"), text.index("Earlier Research"))
                citation = text[text.index("Grounded Resume Systems"):text.index("Earlier Research")]
                self.assertLess(citation.index("J. Example"), citation.index("张三"))
                word = word_text(template_docx(profile, template))
                for value in ("Grounded Resume Systems", "Earlier Research", "J. Example", "张三",
                              "10.1234/resume.1", "Volume 12, pages 1-9"):
                    self.assertIn(value, word)

    def test_publications_only_pdf_passes_text_validation(self):
        profile = ResumeProfile.model_validate({"contact": {"name": "Researcher"}, "publications": PAPERS})
        for template in ("simple", "classic-044"):
            with self.subTest(template=template):
                content, _ = template_pdf(profile, template)
                self.assertTrue(content.startswith(b"%PDF-"))


if __name__ == "__main__":
    unittest.main()
