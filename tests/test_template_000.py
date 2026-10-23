import builtins
import os
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from reportlab.platypus import Paragraph, Table

from resume_app.models import ResumeProfile
from resume_app.pdf import DEFAULT_TEMPLATE_ID, get_template, template_catalog, template_docx, template_pdf
from resume_app.pdf.rendering import load_fonts, normalized_text
from resume_app.pdf.templates.serif_000 import ArrowBullet, Serif000Template
from test_product import ProductTestCase, SOURCE


def complete_profile():
    return ResumeProfile.model_validate({
        **SOURCE.model_dump(),
        "contact": {
            "name": "Sample Applicant", "email": "applicant@example.test", "phone": "+1 555 0100",
            "location": "Remote", "linkedin": "https://linkedin.example.test/sample",
            "website": "https://portfolio.example.test",
        },
        "education": [{
            "institution": "Sample University", "degree": "PhD", "field": "Computer Science",
            "start_date": "2016", "end_date": "2020", "details": ["Education detail"],
        }],
        "publications": [{
            "title": "Verified Publication", "authors": ["First Author", "Second Author"],
            "venue": "Sample Conference", "date": "2025", "doi": "10.1234/example",
            "url": "https://example.test/paper", "details": ["Volume 2, pages 1-8"],
        }],
        "projects": [{"name": "Sample Project", "url": "https://example.test/project", "bullets": ["Project result"]}],
        "certifications": [{"name": "Sample Certificate", "issuer": "Sample Issuer", "date": "2024"}],
        "languages": [{"name": "中文", "proficiency": "母语"}],
    })


def extracted(content):
    return normalized_text("\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages))


class Template000Tests(unittest.TestCase):
    def test_catalog_adds_000_without_replacing_existing_templates_or_default(self):
        catalog = template_catalog()
        self.assertEqual(DEFAULT_TEMPLATE_ID, "classic-044")
        self.assertEqual({item["id"] for item in catalog["items"]}, {"classic-044", "serif-000", "simple"})
        self.assertEqual(get_template("serif-000").reference_document, "references/000.docx")
        self.assertEqual(get_template("classic-044").version, "reportlab-044-v1")
        self.assertEqual(get_template("simple").template_hash, "reportlab-v1")

    def test_measured_source_geometry_and_typography_are_used(self):
        template = get_template("serif-000")
        story = template.story(complete_profile(), load_fonts(serif=True), "")
        self.assertIsInstance(story[0], Table)
        self.assertEqual(sum(isinstance(item, Table) for item in story), 2)
        sections = [item for item in story if isinstance(item, Paragraph) and item.style.name == "Section000"]
        self.assertEqual([item.getPlainText() for item in sections[:3]], ["Industry Experience", "Education", "Publications"])
        self.assertTrue(all(item.style.fontName == "Times-Bold" and item.style.fontSize == 15 and item.style.leading == 19 for item in sections))
        self.assertTrue(all(item.style.backColor is None for item in sections))
        self.assertTrue(all("<u>" in item.text for item in sections))
        self.assertEqual(template.body_size, 11.5)
        self.assertEqual(template.publication_size, 10)
        content, metadata = template_pdf(complete_profile(), "serif-000")
        page = PdfReader(BytesIO(content)).pages[0]
        self.assertAlmostEqual(float(page.mediabox.width), 595.28, delta=0.1)
        self.assertAlmostEqual(float(page.mediabox.height), 841.89, delta=0.1)
        self.assertEqual(metadata, {"template_id": "serif-000", "template_hash": template.template_hash})

    def test_pdf_and_docx_preserve_every_fact_and_order(self):
        profile = complete_profile()
        profile.experience[0].bullets = ["Most important edited point", "Earlier point"]
        content, _ = template_pdf(profile, "serif-000", job_title="Unused target label")
        text = extracted(content)
        for value in (
            *profile.contact.model_dump().values(), "PhD", "Computer Science", "Sample University",
            "2016", "2020", "Education detail", "First Author", "Second Author", "Verified Publication",
            "Sample Conference", "2025", "10.1234/example", "https://example.test/paper",
            "Volume 2, pages 1-8", "Sample Project", "Project result", "Sample Certificate",
            "Sample Issuer", "2024", "中文", "母语",
        ):
            self.assertIn(normalized_text(value), text)
        self.assertLess(text.index("Mostimportanteditedpoint"), text.index("Earlierpoint"))
        self.assertLess(text.index("FirstAuthor"), text.index("SecondAuthor"))
        self.assertNotIn("TARGET:", text)
        document = Document(BytesIO(template_docx(profile, "serif-000")))
        word = normalized_text("\n".join(node.text or "" for node in document.element.iter(qn("w:t"))))
        for value in ("Most important edited point", "First Author", "Volume 2, pages 1-8", "中文", "Sample Applicant"):
            self.assertIn(normalized_text(value), word)
        section = document.sections[0]
        self.assertEqual((section.top_margin.pt, section.bottom_margin.pt), (72, 72))
        self.assertEqual((section.left_margin.pt, section.right_margin.pt), (54, 54))
        self.assertEqual(document.styles["Resume000 Body"].font.name, "Garamond")
        self.assertEqual(document.styles["Resume000 Body"].font.size.pt, 11.5)
        self.assertEqual(document.styles["Resume000 Section"].font.name, "Times New Roman")
        self.assertTrue(document.styles["Resume000 Section"].font.underline)
        self.assertEqual(document.styles["Resume000 Citation"].font.size.pt, 10)

    def test_company_and_dates_share_a_baseline_with_right_aligned_dates(self):
        profile = complete_profile()
        profile.experience[0].company = "Baseline Company"
        profile.experience[0].start_date = "January 2020"
        profile.experience[0].end_date = "December 2024"
        content, _ = template_pdf(profile, "serif-000")
        spans = []

        def visitor(text, cm, tm, _font, _size):
            if text.strip():
                spans.append((
                    text.strip(), tm[4] * cm[0] + tm[5] * cm[2] + cm[4],
                    tm[4] * cm[1] + tm[5] * cm[3] + cm[5],
                ))

        PdfReader(BytesIO(content)).pages[0].extract_text(visitor_text=visitor)
        company = next(item for item in spans if "Baseline Company" in item[0])
        dates = next(item for item in spans if "January 2020" in item[0])
        self.assertAlmostEqual(company[2], dates[2], delta=0.2)
        self.assertAlmostEqual(company[1], 54, delta=0.2)
        self.assertGreater(dates[1], company[1] + 250)
        document = Document(BytesIO(template_docx(profile, "serif-000")))
        entry = document.tables[1]
        self.assertIn("Baseline Company", entry.cell(0, 0).text)
        self.assertEqual(entry.cell(0, 1).text, "January 2020 - December 2024")
        self.assertEqual(entry.cell(0, 1).paragraphs[0].style.paragraph_format.alignment, 2)

    def test_roles_use_sans_serif_italic_and_word_bullets_restore_original_arrow_numbering(self):
        template = get_template("serif-000")
        fonts = load_fonts(serif=True)
        story = template.story(complete_profile(), fonts, "")
        role = next(item for item in story if isinstance(item, Paragraph) and item.style.name == "Role000")
        self.assertNotIn(fonts.italic.fontName if fonts.italic else fonts.regular.fontName, role.text)
        document = Document(BytesIO(template_docx(complete_profile(), "serif-000")))
        self.assertEqual(document.styles["Resume000 Role"].font.name, "Calibri")
        self.assertTrue(document.styles["Resume000 Role"].font.italic)
        self.assertEqual(document.styles["Resume000 Bullet"].paragraph_format.left_indent.pt, 21)
        bullet = next(item for item in document.paragraphs if item.style.name == "Resume000 Bullet")
        number_id = bullet._p.pPr.numPr.numId.val
        numbering = document.part.numbering_part.element
        number = next(node for node in numbering.findall(qn("w:num")) if node.get(qn("w:numId")) == str(number_id))
        abstract_id = number.find(qn("w:abstractNumId")).get(qn("w:val"))
        abstract = next(node for node in numbering.findall(qn("w:abstractNum")) if node.get(qn("w:abstractNumId")) == abstract_id)
        level = abstract.find(qn("w:lvl"))
        self.assertEqual(level.find(qn("w:lvlText")).get(qn("w:val")), "\uf0d8")
        self.assertEqual(level.find(qn("w:rPr")).find(qn("w:rFonts")).get(qn("w:ascii")), "Wingdings")

    def test_arrow_markers_are_font_independent_and_not_repeated_on_continuation_pages(self):
        profile = complete_profile()
        profile.experience[0].bullets = ["Long point " * 700 + "End of long point"]
        profile.education[0].details = []
        profile.projects = []
        content, _ = template_pdf(profile, "serif-000")
        reader = PdfReader(BytesIO(content))
        self.assertGreater(len(reader.pages), 1)
        count = 0
        for page in reader.pages:
            operations = page.get_contents().operations
            for index, (args, operator) in enumerate(operations):
                if operator == b"m" and list(args) == [0, 0]:
                    expected = [([8, 3.5], b"l"), ([0, 7], b"l"), ([2.5, 3.5], b"l"), ([], b"h")]
                    if [(list(values), op) for values, op in operations[index + 1:index + 5]] == expected:
                        count += 1
        self.assertEqual(count, 1)
        self.assertIn("Endoflongpoint", extracted(content))
        template = get_template("serif-000")
        arrows = [item for item in template.story(profile, load_fonts(serif=True), "") if isinstance(item, ArrowBullet)]
        self.assertEqual(len(arrows), 1)
        self.assertEqual(arrows[0].style.leftIndent, 21)

    def test_long_company_dates_and_role_wrap_without_dropping_any_fields(self):
        profile = complete_profile()
        profile.experience[0].company = "Company " * 60
        profile.experience[0].location = "Location " * 45
        profile.experience[0].title = "Senior Engineer " * 25
        profile.experience[0].start_date = "Start date " * 40
        profile.experience[0].end_date = "End date " * 40
        content, _ = template_pdf(profile, "serif-000")
        text = extracted(content)
        for field in ("company", "location", "title", "start_date", "end_date"):
            self.assertIn(normalized_text(getattr(profile.experience[0], field)), text)

    def test_new_template_version_changes_only_000_fingerprint(self):
        template = get_template("serif-000")
        with patch.object(Serif000Template, "version", "reportlab-000-v1"):
            previous = template.template_hash
        self.assertNotEqual(template.template_hash, previous)
        self.assertEqual(get_template("classic-044").version, "reportlab-044-v1")
        self.assertEqual(get_template("simple").template_hash, "reportlab-v1")

    def test_long_bullets_and_unicode_names_continue_without_clipping(self):
        profile = complete_profile()
        profile.contact.name = "张三 " + "A" * 490
        profile.experience[0].bullets = [
            f"Ordered point {index:02d}: " + "reliable services " * 20 for index in range(40)
        ] + ["Long continuous detail " * 400 + "Final sentinel"]
        content, _ = template_pdf(profile, "serif-000")
        self.assertGreater(len(PdfReader(BytesIO(content)).pages), 2)
        text = extracted(content)
        self.assertIn(normalized_text(profile.contact.name), text)
        self.assertIn("Finalsentinel", text)
        self.assertLess(text.index("Orderedpoint00"), text.index("Orderedpoint39"))

    def test_rendering_never_reads_reference_or_calls_office(self):
        original_open = builtins.open

        def no_reference(file, *args, **kwargs):
            if isinstance(file, (str, os.PathLike)) and str(file).lower().endswith(".docx"):
                raise AssertionError("Runtime must not open reference DOCX")
            return original_open(file, *args, **kwargs)

        with (
            patch("builtins.open", side_effect=no_reference),
            patch("subprocess.run", side_effect=AssertionError("Office must not run")),
            patch.object(Serif000Template, "reference_document", "missing.docx"),
        ):
            self.assertTrue(template_pdf(SOURCE, "serif-000")[0].startswith(b"%PDF-"))
        with (
            patch("subprocess.run", side_effect=AssertionError("Office must not run")),
            patch.object(Serif000Template, "reference_document", "missing.docx"),
            patch("resume_app.pdf.templates.serif_000.Document", wraps=Document) as blank_document,
        ):
            self.assertTrue(template_docx(SOURCE, "serif-000").startswith(b"PK"))
            blank_document.assert_called_once_with()

    def test_reference_is_synthetic_and_not_used_in_runtime_hash(self):
        reference = Path(__file__).resolve().parents[1] / "resume_app" / "pdf" / "references" / "000.docx"
        document = Document(reference)
        self.assertEqual(document.core_properties.author, "Sample Applicant")
        self.assertEqual(document.core_properties.last_modified_by, "")
        with ZipFile(reference) as archive:
            content = b"\n".join(archive.read(name) for name in archive.namelist() if name.endswith((".xml", ".rels")))
        self.assertIn(b"Sample Applicant", content)
        regenerated = Document(BytesIO(template_docx(complete_profile(), "serif-000")))
        self.assertEqual(
            [node.text for node in document.element.iter(qn("w:t"))],
            [node.text for node in regenerated.element.iter(qn("w:t"))],
        )
        for forbidden in (b"SynologyDrive", b"Resume_latest"):
            self.assertNotIn(forbidden, content)
        before = get_template("serif-000").template_hash
        with patch.object(Serif000Template, "reference_document", "other.docx"):
            self.assertEqual(get_template("serif-000").template_hash, before)


class Template000ApiTests(ProductTestCase):
    def test_template_can_be_selected_and_downloaded_without_an_llm_call(self):
        self.setup_user()
        self.database.tailored_resumes.insert_one({
            "_id": "template000", "user_id": "local", "job_id": "local:123",
            "job": {"linkedin_id": "123", "title": "Engineer"},
            "resume": complete_profile().model_dump(), "review_status": "pending", "edit_token": "000-token",
        })
        with patch.object(self.workflow, "_complete", side_effect=AssertionError("No model call")):
            response = self.post("/api/tailored/template000/review", {
                "confirmed": True, "edit_token": "000-token", "template_id": "serif-000",
            })
        self.assertEqual(response.status_code, 200, response.get_json())
        detail = self.client.get("/api/tailored/template000").get_json()
        self.assertTrue(detail["pdf_ready"])
        self.assertEqual(detail["template_id"], "serif-000")
        pdf = self.client.get("/api/tailored/template000/pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertIn("VerifiedPublication", extracted(pdf.data))
        word = self.client.get("/api/tailored/template000/download")
        self.assertEqual(word.status_code, 200)
        self.assertEqual(Document(BytesIO(word.data)).styles["Resume000 Body"].font.name, "Garamond")
