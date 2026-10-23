import builtins
import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from reportlab.platypus import Paragraph, Table

from resume_app.documents import DocumentError
from resume_app.models import ResumeProfile
from resume_app.pdf import (
    DEFAULT_TEMPLATE_ID, get_template, template_catalog, template_docx, template_pdf,
)
from resume_app.pdf.rendering import load_fonts, normalized_text
from resume_app.pdf.storage import pdf_files, store_pdf
from resume_app.pdf.templates.classic_044 import Classic044Template
from resume_app.pdf.templates.simple import SimpleTemplate
from resume_app.storage import fingerprint
from test_product import ProductTestCase, SOURCE


def pdf_text(content):
    return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)


def document_text(content):
    document = Document(BytesIO(content))
    return "\n".join(element.text or "" for element in document.element.iter(qn("w:t")))


class ClassTemplateTests(unittest.TestCase):
    def test_classic_preserves_design_and_all_fields_without_sample_content(self):
        profile = SOURCE.model_copy(deep=True)
        profile.contact.name = "张三 <Example> & Resume"
        profile.contact.phone = "+1 555 0100"
        profile.contact.location = "Seattle"
        profile.contact.linkedin = "https://linkedin.example.test/profile"
        profile.contact.website = "https://example.test"
        profile.summary = "软件工程师\nLiteral {{heading}} & <markup>"
        profile.experience.append(profile.experience[0].model_copy(deep=True))
        profile.experience[1].bullets = ["First edited point", "Second edited point"]
        profile.experience[1].location = "Remote"
        profile.education[0].field = "Computer Science"
        profile.education[0].details = ["Education detail"]
        profile = ResumeProfile.model_validate({
            **profile.model_dump(),
            "projects": [{"name": "Project", "url": "https://project.test", "bullets": ["Project result"]}],
            "certifications": [{"name": "Certificate", "issuer": "Issuer", "date": "2026"}],
            "languages": [{"name": "中文", "proficiency": "母语"}],
        })
        content, metadata = template_pdf(profile, job_title="Senior Engineer")
        text = normalized_text(pdf_text(content))
        for value in (
            profile.contact.name, profile.contact.phone, profile.contact.location,
            profile.contact.linkedin, profile.contact.website, "{{heading}} & <markup>",
            "Remote", "Computer Science", "Education detail", "https://project.test",
            "Project result", "Certificate", "Issuer", "2026", "中文", "母语", "TARGET: Senior Engineer",
        ):
            self.assertIn(normalized_text(value), text)
        self.assertLess(text.index("Firsteditedpoint"), text.index("Secondeditedpoint"))
        self.assertNotIn("jianlimoban", text)
        self.assertNotIn("BankofChina", text)
        self.assertEqual(metadata["template_id"], DEFAULT_TEMPLATE_ID)
        page = PdfReader(BytesIO(content)).pages[0]
        self.assertAlmostEqual(float(page.mediabox.width), 595.28, delta=0.1)
        template = get_template(DEFAULT_TEMPLATE_ID)
        story = template.story(profile, load_fonts(serif=True), "Senior Engineer")
        sections = [item for item in story if isinstance(item, Paragraph) and item.style.backColor]
        self.assertEqual(len(sections), 7)
        self.assertEqual(sections[0].style.backColor.hexval(), "0xd9d9d9")
        tables = [item for item in story if isinstance(item, Table)]
        self.assertEqual(len(tables), 6)
        self.assertEqual(story[0].style.fontSize, 28)
        fonts = page["/Resources"]["/Font"].get_object().values()
        self.assertTrue(any(
            "/FontFile2" in font.get_object().get("/FontDescriptor", {})
            for font in fonts
        ))

    def test_pdf_never_opens_reference_word_or_starts_office(self):
        original_open = builtins.open

        def no_docx(file, *args, **kwargs):
            if isinstance(file, (str, os.PathLike)) and str(file).lower().endswith(".docx"):
                raise AssertionError("PDF rendering must not read any Word document")
            return original_open(file, *args, **kwargs)

        with (
            patch("builtins.open", side_effect=no_docx),
            patch("subprocess.run", side_effect=AssertionError("Office must not run")),
            patch.object(Classic044Template, "reference_document", "references/missing.docx"),
            patch("resume_app.pdf.templates.classic_044.Document", side_effect=AssertionError("No DOCX generation")),
        ):
            content, _ = template_pdf(SOURCE)
        self.assertIn(SOURCE.contact.name, pdf_text(content))

    def test_reference_asset_is_present_but_not_part_of_runtime_fingerprint(self):
        template = get_template(DEFAULT_TEMPLATE_ID)
        reference = Path(__file__).resolve().parents[1] / "resume_app" / "pdf" / template.reference_document
        self.assertTrue(reference.is_file())
        before = template.template_hash
        with patch.object(Classic044Template, "reference_document", "another-reference.docx"):
            self.assertEqual(template.template_hash, before)
        with patch.object(Classic044Template, "version", "reportlab-044-v2"):
            self.assertNotEqual(template.template_hash, before)

    def test_empty_sections_contacts_and_bullets_are_omitted(self):
        profile = ResumeProfile.model_validate({
            "contact": {"name": "Only Name"},
            "experience": [{"title": "Engineer", "bullets": ["", "Actual point", ""]}],
        })
        content, _ = template_pdf(profile)
        text = pdf_text(content)
        for omitted in ("PROFILE", "EDUCATION", "SKILLS", "PROJECTS", "CERTIFICATIONS", "LANGUAGES", "TARGET:"):
            self.assertNotIn(omitted, text)
        self.assertIn("Actual point", text)
        self.assertEqual(len(PdfReader(BytesIO(content)).pages), 1)

    def test_long_experience_and_single_long_bullet_continue_without_clipping(self):
        profile = SOURCE.model_copy(deep=True)
        profile.contact.name = "张三 Example"
        profile.summary = "软件工程师 & Python <literal markup>"
        profile.experience[0].bullets = [
            f"Verified point {index:03d}: built reliable Python services and improved system performance."
            for index in range(90)
        ]
        profile.experience[0].bullets.append("Long detail " * 800 + "Final detail sentinel")
        content, _ = template_pdf(profile, job_title="Software Engineer")
        reader = PdfReader(BytesIO(content))
        self.assertGreater(len(reader.pages), 2)
        text = normalized_text(pdf_text(content))
        self.assertIn("张三", text)
        self.assertLess(text.index("Verifiedpoint000"), text.index("Verifiedpoint089"))
        self.assertIn("Finaldetailsentinel", text)
        self.assertIn("ExampleUniversity", text)

    def test_pdf_text_limit_and_maximum_name_are_enforced_without_truncation(self):
        long_name = ResumeProfile.model_validate({"contact": {"name": "A" * 500}, "summary": "Summary"})
        content, _ = template_pdf(long_name)
        self.assertIn("A" * 500, normalized_text(pdf_text(content)))
        profile = ResumeProfile.model_validate({
            "contact": {"name": "A"},
            "experience": [{"bullets": ["x" * 10000] * 7 + ["x" * 9999]}],
        })
        self.assertTrue(template_pdf(profile)[0].startswith(b"%PDF-"))
        profile.experience[0].bullets[-1] += "x"
        with self.assertRaisesRegex(DocumentError, "80,000"):
            template_pdf(profile)

    def test_invalid_templates_fonts_profiles_and_job_titles_are_explicit_errors(self):
        with self.assertRaises(DocumentError):
            get_template("../../private.docx")
        with self.assertRaises(DocumentError):
            template_pdf(ResumeProfile())
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(DocumentError, "字体不存在"):
                template_pdf(SOURCE, font_path=Path(directory) / "missing.ttf")
        for value in ("Invalid\x00control", "Invalid\uffff", "\U0010ffff"):
            invalid = SOURCE.model_copy(deep=True)
            invalid.summary = value
            with self.subTest(value=value), self.assertRaises(DocumentError):
                template_pdf(invalid)
        with self.assertRaisesRegex(DocumentError, "无效控制字符"):
            template_pdf(SOURCE, job_title="Invalid\x00title")

    def test_incomplete_renderer_output_is_not_reported_as_success(self):
        truncated = SimpleTemplate().render(ResumeProfile.model_validate({
            "contact": {"name": SOURCE.contact.name}, "summary": "Truncated resume",
        }))
        with patch.object(Classic044Template, "render", return_value=truncated):
            with self.assertRaisesRegex(DocumentError, "文字缺失"):
                template_pdf(SOURCE)

    def test_docx_compatibility_export_is_generated_by_the_class_without_reading_reference(self):
        with patch.object(Classic044Template, "reference_document", "missing.docx"):
            result = template_docx(SOURCE, job_title="Engineer")
        text = document_text(result)
        self.assertIn(SOURCE.contact.name, text)
        self.assertIn(SOURCE.experience[0].bullets[0], text)
        document = Document(BytesIO(result))
        self.assertEqual(document.paragraphs[0].style.style_id, "ResumeName")
        self.assertEqual(document.styles["Resume Name"].font.size.pt, 28)
        self.assertEqual(document.sections[0].page_width.twips, 11906)
        self.assertTrue(document.tables)


class TemplateApiTests(ProductTestCase):
    def setUp(self):
        super().setUp()
        self.setup_user()
        self.database.tailored_resumes.insert_one({
            "_id": "template-resume", "user_id": "local", "job_id": "local:123",
            "job": {"linkedin_id": "123", "title": "Engineer"},
            "resume": SOURCE.model_dump(), "review_status": "pending", "edit_token": "first",
        })

    def detail(self):
        return self.client.get("/api/tailored/template-resume").get_json()

    def approve(self, template_id=None):
        body = {"confirmed": True, "edit_token": self.detail()["edit_token"]}
        if template_id is not None:
            body["template_id"] = template_id
        return self.post("/api/tailored/template-resume/review", body)

    def test_catalog_has_default_and_empty_previews_without_exposing_reference_paths(self):
        response = self.client.get("/api/resume-templates")
        self.assertEqual(response.get_json(), template_catalog())
        self.assertEqual(response.get_json()["default_template_id"], DEFAULT_TEMPLATE_ID)
        self.assertTrue(all(item["preview_url"] is None for item in response.get_json()["items"]))
        self.assertNotIn("reference_document", response.get_json()["items"][0])
        self.client.post("/api/auth/logout", headers={"X-CSRF-Token": self.token()})
        self.assertEqual(self.client.get("/api/resume-templates").status_code, 401)

    def test_default_class_generation_switching_storage_and_compatibility_download(self):
        with patch("subprocess.run", side_effect=AssertionError("Office must not run")):
            response = self.approve()
        self.assertEqual(response.status_code, 200, response.get_json())
        first = self.detail()
        self.assertEqual(first["template_id"], DEFAULT_TEMPLATE_ID)
        self.assertTrue(first["pdf_ready"])
        word = self.client.get("/api/tailored/template-resume/download")
        self.assertEqual(word.status_code, 200)
        self.assertEqual(Document(BytesIO(word.data)).paragraphs[0].style.style_id, "ResumeName")
        file = pdf_files(self.database).get(first["pdf"]["file_id"])
        self.assertEqual(file.metadata["template_hash"], first["pdf"]["template_hash"])
        self.assertEqual(self.approve("simple").status_code, 200)
        second = self.detail()
        self.assertEqual(second["template_id"], "simple")
        self.assertNotEqual(second["pdf"]["file_id"], first["pdf"]["file_id"])
        self.assertTrue(second["pdf_ready"])
        word = self.client.get("/api/tailored/template-resume/download")
        self.assertNotEqual(Document(BytesIO(word.data)).paragraphs[0].style.style_id, "ResumeName")

    def test_invalid_template_and_render_failure_preserve_the_confirmed_artifact(self):
        self.assertEqual(self.approve("simple").status_code, 200)
        first = self.detail()
        content = self.client.get("/api/tailored/template-resume/pdf").data
        self.assertEqual(self.approve("../template.docx").status_code, 400)
        with patch.object(Classic044Template, "render", side_effect=DocumentError("Synthetic render failure")):
            response = self.approve(DEFAULT_TEMPLATE_ID)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.detail()["pdf"], first["pdf"])
        self.assertEqual(self.client.get("/api/tailored/template-resume/pdf").data, content)
        self.database.tailored_resumes.update_one(
            {"_id": "template-resume"}, {"$set": {"template_id": DEFAULT_TEMPLATE_ID}},
        )
        self.assertFalse(self.detail()["pdf_ready"])
        self.assertEqual(self.client.get("/api/tailored/template-resume/pdf").status_code, 409)

    def test_existing_word_generated_and_legacy_artifacts_do_not_require_regeneration(self):
        content = SimpleTemplate().render(SOURCE)
        for template in (None, {"template_id": DEFAULT_TEMPLATE_ID, "template_hash": "old-word-file-hash"}):
            with self.subTest(template=template):
                artifact = store_pdf(
                    self.database, "local", "template-resume", "saved-file",
                    fingerprint(SOURCE.model_dump()), content, template=template,
                )
                update = {"$set": {"pdf": artifact, "review_status": "reviewed"}}
                if template:
                    update["$set"]["template_id"] = DEFAULT_TEMPLATE_ID
                self.database.tailored_resumes.update_one({"_id": "template-resume"}, update)
                self.assertTrue(self.detail()["pdf_ready"])
                with patch.object(Classic044Template, "render", side_effect=AssertionError("No rerender")):
                    response = self.client.get("/api/tailored/template-resume/pdf")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.data, content)
                pdf_files(self.database).delete("saved-file")


if __name__ == "__main__":
    unittest.main()
