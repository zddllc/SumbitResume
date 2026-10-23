import json
import tempfile
import time
import unittest
from io import BytesIO
from pathlib import Path
from threading import Event
from unittest.mock import Mock, patch
from zipfile import ZipFile

import httpx
import mongomock
from mongomock.gridfs import enable_gridfs_integration
from docx import Document
from openai import AuthenticationError
from pydantic import ValidationError
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject, DictionaryObject, NameObject,
)

from resume_app import create_app
from resume_app.documents import DocumentError, extract_text, resume_docx
from resume_app.errors import AppError
from resume_app.linkedin import Candidate, LinkedInBrowser, RECOMMENDED_URL, job_id_from_url
from resume_app.llm_workflow import apply_plan, parse_output
from resume_app.models import JobPost, LLMSettings, ResumeProfile, TailoringPlan
from resume_app.storage import SecretStore, fingerprint, initialize_database
from run import single_instance

enable_gridfs_integration()


SOURCE = ResumeProfile.model_validate({
    "contact": {"name": "Sample Applicant", "email": "sample@example.test"},
    "summary": "Python engineer building reliable services.",
    "skills": ["Python", "SQL"],
    "experience": [{
        "company": "Example Labs", "title": "Engineer",
        "start_date": "2020", "end_date": "2024",
        "bullets": ["Built Python services with 20% lower latency."],
    }],
    "education": [{"institution": "Example University", "degree": "BSc"}],
})
PLAN = {
    "summary": "Python engineer focused on reliable services.",
    "skills": ["Python"],
    "experience": [{"source_index": 0, "bullets": [{
        "source_index": 0, "text": "Built Python services, lowering latency by 20%.",
    }]}],
    "projects": [],
}
JOB = JobPost(
    linkedin_id="123456", url="https://www.linkedin.com/jobs/view/123456/",
    title="Python Engineer", company="Sample Hiring",
    description="Build reliable Python services and collaborate with the engineering team.",
)


def word_bytes(text="Sample Applicant\nPython engineer"):
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "SQL skills"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def candidate_stream(items):
    yield from items


class ProductTestCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.database = mongomock.MongoClient(tz_aware=True).resume
        self.app = create_app({
            "TESTING": True, "INSTANCE_DIR": Path(self.temp.name) / "instance",
            "BROWSER_PROFILE_ROOT": Path(self.temp.name) / "browser",
            "APP_BASE_PATH": "",
        }, database=self.database)
        self.runner = self.app.extensions["resume_runner"]
        self.addCleanup(self.runner.close)
        self.workflow = self.app.extensions["resume_workflow"]
        self.vault = self.app.extensions["resume_secrets"]
        self.client = self.app.test_client()
        self.client.get("/api/auth/session")

    def token(self):
        with self.client.session_transaction() as session:
            return session["csrf"]

    def post(self, path, body=None, **kwargs):
        headers = {"X-CSRF-Token": self.token()}
        if body is not None:
            kwargs["json"] = body
        response = self.client.post(path, headers=headers, **kwargs)
        self.addCleanup(response.close)
        self.addCleanup(response.request.environ["wsgi.input"].close)
        return response

    def setup_user(self):
        response = self.post("/api/auth/setup", {
            "name": "Sample User", "email": "owner@example.test",
            "password": "local-test-password-123",
        })
        self.assertEqual(response.status_code, 201)

    def settings(self):
        response = self.post("/api/llm", {
            "provider": "deepseek", "model": "deepseek-flash",
            "base_url": "", "api_key": "synthetic-api-key",
        })
        self.assertEqual(response.status_code, 200, response.get_json())

    def confirmed(self):
        self.database.resumes.update_one(
            {"_id": "local"},
            {"$set": {"draft": SOURCE.model_dump(), "edit_token": "first-draft"}},
            upsert=True,
        )
        self.workflow.confirm_resume("local", "first-draft", SOURCE)

    def wait_task(self, task_id):
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            task = self.database.tasks.find_one({"_id": task_id})
            if task and task["status"] not in {"running", "queued", "cancelling"}:
                return task
            time.sleep(.01)
        self.fail("Background test task did not complete")


class AuthenticationTests(ProductTestCase):
    def test_session_is_json_and_backend_no_longer_serves_frontend_assets(self):
        response = self.client.get("/api/auth/session")
        self.assertEqual(response.mimetype, "application/json")
        self.assertTrue(response.get_json()["setup_required"])
        self.assertIsNone(response.get_json()["user"])
        self.assertTrue(response.get_json()["csrf_token"])
        self.assertEqual(self.client.get("/static/app.js").status_code, 404)

    def test_setup_is_single_user_and_hashes_password(self):
        self.setup_user()
        user = self.database.users.find_one({"_id": "local"})
        self.assertTrue(user["password_hash"].startswith("scrypt:"))
        self.assertNotIn("local-test-password-123", user["password_hash"])
        self.assertEqual(self.post("/api/auth/setup", {}).status_code, 403)
        self.assertEqual(self.database.users.count_documents({}), 1)

    def test_setup_rejects_invalid_email_and_short_password(self):
        response = self.post("/api/auth/setup", {
            "name": "Sample", "email": "bad-email", "password": "short",
        })
        self.assertEqual(response.status_code, 400)
        response = self.post("/api/auth/setup", {
            "name": "Sample", "email": "a@example.test", "password": "short",
        })
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.database.users.count_documents({}), 0)

    def test_missing_csrf_is_rejected_even_for_setup(self):
        response = self.client.post("/api/auth/setup", json={
            "name": "Sample", "email": "a@example.test", "password": "long-password-123",
        })
        self.assertEqual(response.status_code, 403)

    def test_password_minimum_is_eight_characters_in_api_setup_and_login(self):
        account = {"name": "Sample", "email": "a@example.test"}
        rejected = self.post("/api/auth/setup", {**account, "password": "1234567"})
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(self.database.users.count_documents({}), 0)
        accepted = self.post("/api/auth/setup", {**account, "password": "12345678"})
        self.assertEqual(accepted.status_code, 201)
        self.post("/api/auth/logout")
        self.client.get("/api/auth/session")
        logged_in = self.post("/api/auth/login", {
            "email": account["email"], "password": "12345678",
        })
        self.assertEqual(logged_in.status_code, 200)
        self.assertEqual(self.client.get("/api/state").status_code, 200)

    def test_login_logout_and_session_security(self):
        self.setup_user()
        self.assertEqual(self.client.get("/api/state").status_code, 200)
        self.post("/api/auth/logout")
        self.assertEqual(self.client.get("/api/state").status_code, 401)
        self.client.get("/api/auth/session")
        response = self.post("/api/auth/login", {
            "email": "OWNER@example.test", "password": "local-test-password-123",
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn("HttpOnly", response.headers["Set-Cookie"])
        self.assertIn("SameSite=Strict", response.headers["Set-Cookie"])
        self.assertEqual(self.client.get("/api/state").status_code, 200)

    def test_failed_logins_are_rate_limited(self):
        self.setup_user()
        self.post("/api/auth/logout")
        self.client.get("/api/auth/session")
        for _ in range(5):
            self.assertEqual(self.post("/api/auth/login", {
                "email": "owner@example.test", "password": "wrong",
            }).status_code, 401)
        self.assertEqual(self.post("/api/auth/login", {
            "email": "owner@example.test", "password": "wrong",
        }).status_code, 429)

    def test_json_auth_preserves_password_whitespace(self):
        password = "  pass12  "
        created = self.post("/api/auth/setup", {
            "name": "Whitespace User", "email": "space@example.test", "password": password,
        })
        self.assertEqual(created.status_code, 201)
        self.post("/api/auth/logout")
        self.assertEqual(self.post("/api/auth/login", {
            "email": "space@example.test", "password": password.strip(),
        }).status_code, 401)
        self.assertEqual(self.post("/api/auth/login", {
            "email": "space@example.test", "password": password,
        }).status_code, 200)

    def test_api_auth_csrf_and_response_headers(self):
        self.assertEqual(self.client.get("/api/state").status_code, 401)
        self.setup_user()
        response = self.client.post("/api/personal", json={})
        self.assertEqual(response.status_code, 403)
        response = self.client.post(
            "/api/personal", json={}, headers={"X-CSRF-Token": "non-matching"}
        )
        self.assertEqual(response.status_code, 403)
        response = self.client.get("/api/state")
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertIn("frame-ancestors 'none'", response.headers["Content-Security-Policy"])
        self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
        self.assertEqual(self.client.get("/", headers={"Host": "attacker.test"}).status_code, 400)


class VaultTests(ProductTestCase):
    def setUp(self):
        super().setUp()
        self.setup_user()

    def test_key_is_encrypted_and_not_returned(self):
        self.settings()
        stored = self.database.llm_settings.find_one({"_id": "local"})
        self.assertNotIn("synthetic-api-key", str(stored))
        self.assertEqual(self.workflow.llm_settings("local").api_key, "synthetic-api-key")
        body = self.client.get("/api/state").get_json()
        self.assertTrue(body["llm"]["key_saved"])
        self.assertNotIn("api_key", body["llm"])
        self.assertNotIn("key_encrypted", body["llm"])

    def test_key_cannot_be_reused_for_changed_endpoint(self):
        self.settings()
        response = self.post("/api/llm", {
            "provider": "deepseek", "model": "deepseek-flash",
            "base_url": "https://new-endpoint.example.test/v1",
        })
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.post("/api/llm", {
            "provider": "deepseek", "model": "custom-model",
        }).status_code, 200)
        self.assertEqual(self.workflow.llm_settings("local").api_key, "synthetic-api-key")

    def test_removed_providers_and_insecure_endpoints_are_rejected(self):
        for provider in ("agnes", "proxy"):
            self.assertEqual(self.post("/api/llm", {
                "provider": provider, "api_key": "synthetic-api-key",
            }).status_code, 400)
        for endpoint in (
            "http://example.test/v1", "https://user:password@example.test",
            "https://example.test?key=secret",
        ):
            self.assertEqual(self.post("/api/llm", {
                "provider": "openai", "model": "test", "base_url": endpoint,
                "api_key": "synthetic-api-key",
            }).status_code, 400)

    def test_personal_answers_are_encrypted_and_remain_optional(self):
        response = self.post("/api/personal", {
            "work_authorization": "Sample work permit",
            "veteran_status": "decline", "disability_status": "unspecified",
        })
        self.assertEqual(response.status_code, 200)
        stored = self.database.users.find_one({"_id": "local"})
        self.assertNotIn("Sample work permit", str(stored))
        self.assertNotIn("veteran_status", stored)
        personal = self.client.get("/api/state").get_json()["personal"]
        self.assertEqual(personal["veteran_status"], "decline")
        self.assertEqual(personal["disability_status"], "unspecified")
        self.assertEqual(self.post("/api/personal", {"veteran_status": "inferred"}).status_code, 400)

    def test_external_credentials_are_encrypted_masked_and_owner_scoped(self):
        response = self.post("/api/credentials", {
            "site": "jobs.example.test", "label": "Sample portal",
            "username": "synthetic-user@example.test", "password": "synthetic-password",
        })
        self.assertEqual(response.status_code, 201)
        stored = self.database.credentials.find_one({"user_id": "local"})
        self.assertNotIn("synthetic-password", str(stored))
        self.assertNotIn("synthetic-user", str(stored))
        self.assertEqual(
            self.vault.decrypt(stored["secret_encrypted"])["password"], "synthetic-password"
        )
        self.database.credentials.insert_one({
            "_id": "other-account", "user_id": "someone-else", "site": "https://private.test",
        })
        body = self.client.get("/api/state").get_json()
        self.assertEqual(len(body["credentials"]), 1)
        self.assertEqual(body["credentials"][0]["username_hint"], "sy***")
        self.assertNotIn("password", str(body["credentials"]))
        self.assertEqual(self.client.delete(
            "/api/credentials/other-account", headers={"X-CSRF-Token": self.token()},
        ).status_code, 404)
        self.assertEqual(self.client.delete(
            f"/api/credentials/{stored['_id']}", headers={"X-CSRF-Token": self.token()},
        ).status_code, 200)

    def test_secret_key_is_persistent_and_wrong_key_fails_explicitly(self):
        encrypted = self.vault.encrypt({"value": "synthetic-secret"})
        restored = SecretStore(Path(self.temp.name) / "instance" / "secrets.json")
        self.assertEqual(restored.decrypt(encrypted), {"value": "synthetic-secret"})
        self.assertEqual(restored.session_key, self.vault.session_key)
        other = SecretStore(Path(self.temp.name) / "different" / "secrets.json")
        with self.assertRaisesRegex(RuntimeError, "无法解密"):
            other.decrypt(encrypted)

    def test_external_password_whitespace_is_preserved(self):
        secret = "  meaningful password whitespace  "
        self.assertEqual(self.post("/api/credentials", {
            "site": "jobs.example.test", "username": "sample", "password": secret,
        }).status_code, 201)
        stored = self.database.credentials.find_one({})
        self.assertEqual(self.vault.decrypt(stored["secret_encrypted"])["password"], secret)


class ResumeAndJobTests(ProductTestCase):
    def setUp(self):
        super().setUp()
        self.setup_user()
        self.settings()

    def test_upload_parse_requires_consent_and_confirmation_before_scan(self):
        response = self.post("/api/resume/upload", data={
            "file": (BytesIO(word_bytes()), "resume.docx"),
        })
        self.assertEqual(response.status_code, 400)
        with patch.object(self.workflow, "_complete", return_value=SOURCE.model_dump_json()):
            response = self.post("/api/resume/upload", data={
                "file": (BytesIO(word_bytes()), "resume.docx"), "consent": "yes",
            })
            self.assertEqual(response.status_code, 202, response.get_json())
            task = self.wait_task(response.get_json()["task_id"])
        self.assertEqual(task["status"], "succeeded")
        record = self.database.resumes.find_one({"_id": "local"})
        self.assertEqual(record["draft"]["contact"]["name"], SOURCE.contact.name)
        self.assertNotIn("confirmed", record)
        self.assertEqual(self.post("/api/jobs/scan", {
            "limit": 1, "consent": True,
        }).status_code, 400)
        self.assertEqual(self.post("/api/resume/confirm", {
            "edit_token": record["edit_token"], "profile": SOURCE.model_dump(),
        }).status_code, 200)
        stored = self.database.resumes.find_one({"_id": "local"})
        self.assertIsNone(stored["draft"])
        self.assertEqual(stored["confirmed_hash"], fingerprint(SOURCE.model_dump()))

    def test_upload_file_size_limit_excludes_multipart_overhead(self):
        limit = 10 * 1024 * 1024
        with (
            patch("resume_app.web.extract_text", return_value="Synthetic resume") as extract,
            patch.object(self.runner, "parse_resume", return_value="synthetic-task"),
        ):
            accepted = self.post("/api/resume/upload", data={
                "file": (BytesIO(b"x" * limit), "resume.docx"), "consent": "yes",
            })
            self.assertEqual(accepted.status_code, 202, accepted.get_json())
            self.assertEqual(len(extract.call_args.args[1]), limit)
            rejected = self.post("/api/resume/upload", data={
                "file": (BytesIO(b"x" * (limit + 1)), "resume.docx"), "consent": "yes",
            })
            self.assertEqual(rejected.status_code, 413)
            self.assertEqual(extract.call_count, 1)

    def test_new_draft_does_not_replace_confirmed_source_and_stale_confirmation_fails(self):
        self.confirmed()
        before = self.database.resumes.find_one({"_id": "local"})
        with patch.object(self.workflow, "_complete", return_value=SOURCE.model_dump_json()):
            self.workflow.parse_resume("local", "new.docx", "Sample source", self.workflow.llm_settings("local"))
        self.assertEqual(self.workflow.confirmed_resume("local"), SOURCE)
        response = self.post("/api/resume/confirm", {
            "edit_token": before["edit_token"], "profile": SOURCE.model_dump(),
        })
        self.assertEqual(response.status_code, 409)

    def test_saved_resume_can_be_edited_after_login_without_upload_or_llm(self):
        self.confirmed()
        self.database.llm_settings.delete_one({"_id": "local"})
        self.post("/api/auth/logout")
        self.client.get("/api/auth/session")
        self.assertEqual(self.post("/api/auth/login", {
            "email": "owner@example.test", "password": "local-test-password-123",
        }).status_code, 200)
        record = self.client.get("/api/state").get_json()["resume"]
        self.assertIsNone(record["draft"])
        edited = record["confirmed"]
        bullets = ["Revised first bullet\nwith a continuation.", "New independent bullet."]
        edited["experience"][0]["bullets"] = bullets
        with patch.object(self.workflow, "_complete") as complete:
            response = self.post("/api/resume/confirm", {
                "edit_token": record["edit_token"], "profile": edited,
            })
        self.assertEqual(response.status_code, 200)
        complete.assert_not_called()
        self.post("/api/auth/logout")
        self.client.get("/api/auth/session")
        self.assertEqual(self.post("/api/auth/login", {
            "email": "owner@example.test", "password": "local-test-password-123",
        }).status_code, 200)
        saved = self.client.get("/api/state").get_json()["resume"]
        self.assertEqual(saved["confirmed"]["experience"][0]["bullets"], bullets)
        self.assertNotEqual(saved["edit_token"], record["edit_token"])
        self.assertNotEqual(saved["confirmed_hash"], record["confirmed_hash"])
        exported = Document(resume_docx(ResumeProfile.model_validate(saved["confirmed"])))
        self.assertIn(bullets[0], [paragraph.text for paragraph in exported.paragraphs])

    def test_full_scan_cache_review_and_word_download(self):
        self.confirmed()
        fake_browser = Mock()
        fake_browser.jobs.side_effect = lambda _limit, _cancel: candidate_stream([Candidate(JOB.linkedin_id, post=JOB)])
        self.runner.browser = fake_browser
        with patch.object(self.workflow, "_complete", return_value=json.dumps(PLAN)) as complete:
            for _ in range(2):
                response = self.post("/api/jobs/scan", {"limit": 1, "consent": True})
                self.assertEqual(response.status_code, 202, response.get_json())
                task = self.wait_task(response.get_json()["task_id"])
                self.assertEqual(task["status"], "succeeded", task)
            self.assertEqual(complete.call_count, 1)
        self.assertEqual(task["cached"], 1)
        self.assertEqual(self.database.jobs.count_documents({}), 1)
        self.assertEqual(self.database.tailored_resumes.count_documents({}), 1)
        listing = self.client.get("/api/jobs").get_json()["items"]
        self.assertTrue(listing[0]["current_source"])
        detail = self.client.get(f"/api/jobs/{JOB.linkedin_id}").get_json()
        resume_id = detail["versions"][0]["id"]
        stored = self.client.get(f"/api/tailored/{resume_id}").get_json()
        self.assertEqual(stored["resume"]["experience"][0]["company"], "Example Labs")
        self.assertEqual(stored["job"]["description"], JOB.description)
        self.assertEqual(stored["review_status"], "pending")
        download = self.client.get(f"/api/tailored/{resume_id}/download")
        self.assertEqual(download.status_code, 200)
        generated = Document(BytesIO(download.data))
        self.assertIn(SOURCE.contact.name, "\n".join(p.text for p in generated.paragraphs))
        self.assertEqual(self.post(f"/api/tailored/{resume_id}/review", {
            "confirmed": True, "edit_token": stored["edit_token"],
            "template_id": "simple",
        }).status_code, 200)
        self.assertEqual(self.client.get(f"/api/tailored/{resume_id}").get_json()["review_status"], "reviewed")

    def test_work_point_order_is_saved_and_exported_without_llm(self):
        self.confirmed()
        record = self.client.get("/api/state").get_json()["resume"]
        edited = record["confirmed"]
        ordered = ["Highest priority\nand most recent", "Second priority", "Earlier point"]
        edited["experience"][0]["bullets"] = ordered
        with patch.object(self.workflow, "_complete") as complete:
            response = self.post("/api/resume/confirm", {
                "edit_token": record["edit_token"], "profile": edited,
            })
        self.assertEqual(response.status_code, 200)
        complete.assert_not_called()
        stored = self.client.get("/api/state").get_json()["resume"]["confirmed"]
        self.assertEqual(stored["experience"][0]["bullets"], ordered)
        exported = Document(resume_docx(ResumeProfile.model_validate(stored)))
        self.assertEqual(
            [paragraph.text for paragraph in exported.paragraphs if paragraph.style.name == "List Bullet"],
            ordered,
        )

    def test_reordering_changes_prompt_cache_and_generated_work_point_order(self):
        source = SOURCE.model_copy(deep=True)
        source.experience[0].bullets = ["First point", "Middle point", "Last point"]
        settings = self.workflow.llm_settings("local")

        def complete(_settings, prompt, _system):
            bullets = json.loads(prompt)["confirmed_resume"]["experience"][0]["bullets"]
            return json.dumps({
                **PLAN, "experience": [{"source_index": 0, "bullets": [
                    {"source_index": index, "text": bullets[index]} for index in (2, 0)
                ]}],
            })

        with patch.object(self.workflow, "_complete", side_effect=complete) as call:
            first_id, _ = self.workflow.tailor("local", JOB, source, settings)
            source.experience[0].bullets.reverse()
            second_id, cached = self.workflow.tailor("local", JOB, source, settings)
        self.assertFalse(cached)
        self.assertNotEqual(first_id, second_id)
        self.assertEqual(call.call_count, 2)
        self.assertIn("descending importance and recency", call.call_args.args[2])
        self.assertIn("Do not infer dates from position", call.call_args.args[2])
        saved = self.database.tailored_resumes.find_one({"_id": second_id})
        self.assertEqual(saved["resume"]["experience"][0]["bullets"], ["Last point", "First point"])

    def test_new_priority_prompt_does_not_reuse_previous_version_cache(self):
        settings = self.workflow.llm_settings("local")
        with patch.object(self.workflow, "_complete", return_value=json.dumps(PLAN)) as call:
            with patch("resume_app.llm_workflow.PROMPT_VERSION", 1):
                old_id, _ = self.workflow.tailor("local", JOB, SOURCE, settings)
            new_id, cached = self.workflow.tailor("local", JOB, SOURCE, settings)
        self.assertFalse(cached)
        self.assertNotEqual(new_id, old_id)
        self.assertEqual(call.call_count, 2)

    def test_changed_source_or_description_invalidates_cached_resume(self):
        settings = self.workflow.llm_settings("local")
        with patch.object(self.workflow, "_complete", return_value=json.dumps(PLAN)) as complete:
            first, _ = self.workflow.tailor("local", JOB, SOURCE, settings)
            changed = SOURCE.model_copy(deep=True)
            changed.contact.phone = "555-0100"
            second, cached = self.workflow.tailor("local", JOB, changed, settings)
            self.assertFalse(cached)
            changed_job = JOB.model_copy(update={"description": JOB.description + " Remote role."})
            third, cached = self.workflow.tailor("local", changed_job, changed, settings)
            self.assertFalse(cached)
            self.assertEqual(complete.call_count, 3)
        self.assertEqual(len({first, second, third}), 3)

    def test_personal_answers_and_credentials_never_enter_tailoring_prompt(self):
        self.post("/api/personal", {"work_authorization": "private-authorization-answer"})
        self.post("/api/credentials", {
            "site": "jobs.example.test", "username": "private-username",
            "password": "private-site-password",
        })
        with patch.object(self.workflow, "_complete", return_value=json.dumps(PLAN)) as complete:
            self.workflow.tailor("local", JOB, SOURCE, self.workflow.llm_settings("local"))
        prompt = complete.call_args.args[1]
        for secret in ("private-authorization-answer", "private-username", "private-site-password"):
            self.assertNotIn(secret, prompt)
        self.assertNotIn("api_key", str(self.database.tailored_resumes.find_one({})))

    def test_model_failure_saves_job_but_no_success_shaped_resume(self):
        with patch.object(self.workflow, "_complete", return_value="not-json"):
            with self.assertRaises(AppError):
                self.workflow.tailor("local", JOB, SOURCE, self.workflow.llm_settings("local"))
        self.assertEqual(self.database.jobs.count_documents({}), 1)
        self.assertEqual(self.database.tailored_resumes.count_documents({}), 0)

    def test_changed_description_marks_previous_resume_as_historical(self):
        self.confirmed()
        with patch.object(self.workflow, "_complete", return_value=json.dumps(PLAN)):
            self.workflow.tailor("local", JOB, SOURCE, self.workflow.llm_settings("local"))
        changed = JOB.model_copy(update={"description": JOB.description + " Updated responsibilities."})
        self.workflow.save_job("local", changed)
        listing = self.client.get("/api/jobs").get_json()["items"]
        self.assertFalse(listing[0]["current_source"])

    def test_scan_limits_consent_and_missing_browser(self):
        self.confirmed()
        for value in (0, 101, 1.5, True, "20"):
            self.assertEqual(self.post("/api/jobs/scan", {
                "limit": value, "consent": True,
            }).status_code, 400)
        self.assertEqual(self.post("/api/jobs/scan", {"limit": 1}).status_code, 400)
        for limit in (1, 100):
            response = self.post("/api/jobs/scan", {"limit": limit, "consent": True})
            self.assertEqual(response.status_code, 202)
            task = self.wait_task(response.get_json()["task_id"])
            self.assertEqual(task["status"], "failed")
            self.assertIn("Chrome", task["error"])

    def test_tasks_do_not_run_concurrently_and_cancel_persists_terminal_state(self):
        started = Event()
        def blocking(_task, cancel):
            started.set()
            cancel.wait(3)
        task_id = self.runner.start("local", "test", blocking)
        self.assertTrue(started.wait(2))
        with self.assertRaises(AppError):
            self.runner.start("local", "test", lambda _task, _cancel: None)
        response = self.post(f"/api/tasks/{task_id}/cancel")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.wait_task(task_id)["status"], "cancelled")
        self.assertEqual(self.post(f"/api/tasks/{task_id}/cancel").status_code, 409)
        self.assertEqual(self.database.tasks.find_one({"_id": task_id})["status"], "cancelled")

    def test_provider_auth_failure_aborts_scan_without_repeated_requests(self):
        self.confirmed()
        fake_browser = Mock()
        fake_browser.jobs.side_effect = lambda _limit, _cancel: candidate_stream([
            Candidate(JOB.linkedin_id, post=JOB), Candidate("222", post=JOB),
        ])
        self.runner.browser = fake_browser
        response = httpx.Response(401, request=httpx.Request("POST", "https://example.test"))
        error = AuthenticationError("synthetic auth failure", response=response, body={})
        with patch.object(self.workflow, "_complete", side_effect=error) as complete:
            result = self.post("/api/jobs/scan", {"limit": 2, "consent": True})
            task = self.wait_task(result.get_json()["task_id"])
        self.assertEqual(task["status"], "failed")
        self.assertEqual(complete.call_count, 1)

    def test_incomplete_job_is_visible_as_task_error(self):
        self.confirmed()
        fake_browser = Mock()
        fake_browser.jobs.side_effect = lambda _limit, _cancel: candidate_stream([
            Candidate("123", error="Description unavailable"),
        ])
        self.runner.browser = fake_browser
        response = self.post("/api/jobs/scan", {"limit": 1, "consent": True})
        task = self.wait_task(response.get_json()["task_id"])
        self.assertEqual(task["status"], "failed")
        self.assertEqual(task["failed"], 1)
        self.assertEqual(task["errors"][0]["job_id"], "123")

    def test_restart_marks_unfinished_tasks_interrupted(self):
        self.database.tasks.insert_one({"_id": "orphan", "status": "running"})
        initialize_database(self.database)
        task = self.database.tasks.find_one({"_id": "orphan"})
        self.assertEqual(task["status"], "interrupted")
        self.assertIn("重启", task["error"])


class DocumentAndPlanTests(unittest.TestCase):
    def test_second_backend_instance_is_rejected_before_startup(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "app.lock"
            with single_instance(path):
                with self.assertRaisesRegex(RuntimeError, "重复启动"):
                    with single_instance(path):
                        self.fail("Second process lock should not be acquired")
            with single_instance(path):
                pass

    def test_docx_extracts_paragraphs_and_tables(self):
        text = extract_text("resume.docx", word_bytes())
        self.assertIn("Sample Applicant", text)
        self.assertIn("SQL skills", text)

    def test_pdf_text_extraction(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 50 750 Td (Sample Applicant Python Engineer) Tj ET")
        page[NameObject("/Contents")] = stream
        output = BytesIO()
        writer.write(output)
        self.assertIn("Sample Applicant", extract_text("resume.pdf", output.getvalue()))

    def test_invalid_legacy_empty_encrypted_and_scanned_files_fail(self):
        for filename, content in (
            ("old.doc", b"legacy"), ("bad.docx", b"bad"),
            ("bad.pdf", b"bad"),
        ):
            with self.subTest(filename=filename), self.assertRaises(DocumentError):
                extract_text(filename, content)
        empty = BytesIO()
        Document().save(empty)
        with self.assertRaisesRegex(DocumentError, "未读取到文字"):
            extract_text("empty.docx", empty.getvalue())
        for encrypted in (False, True):
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            if encrypted:
                writer.encrypt("synthetic-password")
            stream = BytesIO()
            writer.write(stream)
            with self.subTest(encrypted=encrypted), self.assertRaises(DocumentError):
                extract_text("resume.pdf", stream.getvalue())
        with self.assertRaises(DocumentError):
            extract_text("resume.docx", word_bytes("x" * 80001))

    def test_resume_text_length_boundary_is_exact(self):
        for length in (80000, 80001):
            document = Document()
            document.add_paragraph("x" * length)
            stream = BytesIO()
            document.save(stream)
            if length == 80000:
                self.assertEqual(len(extract_text("resume.docx", stream.getvalue())), 80000)
            else:
                with self.assertRaises(DocumentError):
                    extract_text("resume.docx", stream.getvalue())

    def test_malformed_docx_xml_returns_document_error(self):
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            archive.writestr("[Content_Types].xml", "<invalid-xml")
        with self.assertRaises(DocumentError):
            extract_text("resume.docx", stream.getvalue())

    def test_common_resume_schema_rejects_private_answers(self):
        with self.assertRaises(ValidationError):
            ResumeProfile.model_validate({"veteran_status": "veteran"})

    def test_model_json_is_strict_with_optional_code_fence(self):
        parsed = parse_output("```json\n" + json.dumps(PLAN) + "\n```", TailoringPlan)
        self.assertEqual(parsed.summary, PLAN["summary"])
        with self.assertRaises(AppError):
            parse_output("Here is JSON: " + json.dumps(PLAN), TailoringPlan)

    def test_tailoring_cannot_change_contact_employers_dates_or_education(self):
        result = apply_plan(SOURCE, TailoringPlan.model_validate(PLAN))
        self.assertEqual(result.contact, SOURCE.contact)
        self.assertEqual(result.experience[0].company, SOURCE.experience[0].company)
        self.assertEqual(result.experience[0].start_date, "2020")
        self.assertEqual(result.education, SOURCE.education)
        self.assertNotEqual(result.experience[0].bullets, SOURCE.experience[0].bullets)
        self.assertEqual(SOURCE.experience[0].bullets[0], "Built Python services with 20% lower latency.")

    def test_invalid_skill_source_indexes_and_invented_fields_are_rejected(self):
        for changes in (
            {"skills": ["Unknown invented skill"]},
            {"skills": ["Python", "Python"]},
            {"experience": [{"source_index": 9, "bullets": []}]},
            {"experience": [{"source_index": 0, "bullets": [{"source_index": 9, "text": "Fake"}]}]},
            {"experience": [{"source_index": 0, "bullets": []}, {"source_index": 0, "bullets": []}]},
        ):
            with self.subTest(changes=changes), self.assertRaises(AppError):
                apply_plan(SOURCE, TailoringPlan.model_validate({**PLAN, **changes}))
        with self.assertRaises(ValidationError):
            TailoringPlan.model_validate({**PLAN, "company": "Invented employer"})
        with self.assertRaises(ValidationError):
            TailoringPlan.model_validate({
                **PLAN, "experience": [{"source_index": True, "bullets": []}],
            })

    def test_selected_work_points_preserve_user_priority_but_projects_keep_existing_order(self):
        source = ResumeProfile.model_validate({
            **SOURCE.model_dump(),
            "experience": [{
                **SOURCE.experience[0].model_dump(),
                "bullets": ["Most important", "Middle", "Least important"],
            }],
            "projects": [{"name": "Project", "bullets": ["Project first", "Project second"]}],
        })
        plan = TailoringPlan.model_validate({
            **PLAN,
            "experience": [{"source_index": 0, "bullets": [
                {"source_index": 2, "text": "Rewritten least important"},
                {"source_index": 0, "text": "Rewritten most important"},
            ]}],
            "projects": [{"source_index": 0, "bullets": [
                {"source_index": 1, "text": "Rewritten project second"},
                {"source_index": 0, "text": "Rewritten project first"},
            ]}],
        })
        result = apply_plan(source, plan)
        self.assertEqual(result.experience[0].bullets, [
            "Rewritten most important", "Rewritten least important",
        ])
        self.assertEqual(result.projects[0].bullets, [
            "Rewritten project second", "Rewritten project first",
        ])
        self.assertEqual(source.experience[0].bullets, ["Most important", "Middle", "Least important"])

    def test_word_export_contains_all_common_sections(self):
        document = Document(resume_docx(SOURCE))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Sample Applicant", text)
        self.assertIn("Example Labs", text)
        self.assertIn("Example University", text)
        self.assertIn("20%", text)

    def test_job_url_validation(self):
        for url in (
            "https://www.linkedin.com/jobs/view/123456/?trackingId=test",
            "/jobs/view/engineer-at-company-123456/",
        ):
            self.assertEqual(job_id_from_url(url), "123456")
        for url in (
            "https://attacker.test/jobs/view/123456/",
            "https://linkedin.com.attacker.test/jobs/view/123456/", "/jobs/not-a-job/",
        ):
            self.assertIsNone(job_id_from_url(url))

    def test_browser_requires_manual_login_and_never_bypasses_checkpoint(self):
        browser = LinkedInBrowser(Path("unused-test-profile"))
        driver = Mock()
        driver.current_url = "https://www.linkedin.com/checkpoint/challenge"
        with self.assertRaisesRegex(AppError, "手动"):
            browser._require_login(driver)
        driver.current_url = "https://www.linkedin.com/login"
        with self.assertRaises(AppError):
            browser._require_login(driver)

    def test_job_description_is_loaded_before_expanding(self):
        browser = LinkedInBrowser(Path("unused-test-profile"))
        driver = Mock()
        driver.current_url = JOB.url
        button = Mock()
        button.text = "Show more"
        button.get_attribute.return_value = "false"
        panel = driver.find_element.return_value
        panel.find_elements.return_value = [button]
        expanded = "This is the complete expanded job description for a Python engineering role."
        with (
            patch.object(browser, "_visible_text", side_effect=[
                "Short job description preview", expanded, "Python Engineer", "Sample", "Remote",
            ]),
            patch.object(browser, "_panel_job_id", return_value=JOB.linkedin_id),
            patch.object(browser, "_read_application", return_value={"application_kind": "unavailable"}),
        ):
            result = browser._read_panel(driver, JOB.linkedin_id)
        self.assertEqual(result.description, expanded)
        button.click.assert_called_once()

    def test_already_expanded_job_description_is_not_collapsed(self):
        browser = LinkedInBrowser(Path("unused-test-profile"))
        driver = Mock()
        driver.current_url = JOB.url
        button = Mock()
        button.text = "Show less"
        button.get_attribute.return_value = "true"
        panel = driver.find_element.return_value
        panel.find_elements.return_value = [button]
        with (
            patch.object(browser, "_visible_text", side_effect=[
                JOB.description, JOB.title, JOB.company, JOB.location,
            ]),
            patch.object(browser, "_panel_job_id", return_value=JOB.linkedin_id),
            patch.object(browser, "_read_application", return_value={"application_kind": "unavailable"}),
        ):
            result = browser._read_panel(driver, JOB.linkedin_id)
        self.assertEqual(result.description, JOB.description)
        button.click.assert_not_called()

    def test_job_traversal_limit_uses_current_split_view_without_extra_job_tabs(self):
        browser = LinkedInBrowser(Path("unused-test-profile"))
        driver = Mock()
        driver.current_window_handle = "listing"
        driver.current_url = RECOMMENDED_URL + "?currentJobId=123456"
        driver.window_handles = ["listing", "details"]
        driver.switch_to.new_window.side_effect = lambda _kind: setattr(
            driver, "current_window_handle", "details"
        )
        driver.switch_to.window.side_effect = lambda handle: setattr(
            driver, "current_window_handle", handle
        )
        browser.driver = driver
        with (
            patch.object(browser, "_require_login"),
            patch.object(browser, "_card_ids", return_value=["123456", "999999"]),
            patch.object(browser, "_read_job", return_value=JOB) as read,
        ):
            results = list(browser.jobs(1, Event()))
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].post, JOB)
        read.assert_called_once_with(driver, "123456")
        driver.get.assert_not_called()
        driver.close.assert_not_called()
        driver.switch_to.new_window.assert_not_called()
        self.assertEqual(driver.current_window_handle, "listing")

    def test_docx_zip_expansion_is_bounded(self):
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            for index in range(2001):
                archive.writestr(f"empty-{index}", "")
        with self.assertRaisesRegex(DocumentError, "过大"):
            extract_text("resume.docx", stream.getvalue())


if __name__ == "__main__":
    unittest.main()
