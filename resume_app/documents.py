"""Text extraction and generated DOCX downloads; source files are not retained."""

from io import BytesIO
from collections.abc import Sequence
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from resume_app.models import Model, Publication, ResumeProfile


MAX_TEXT_LENGTH = 80000
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


class DocumentError(ValueError):
    pass


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension == ".docx":
        try:
            with ZipFile(BytesIO(content)) as archive:
                expanded = sum(item.file_size for item in archive.infolist())
                if expanded > 30 * 1024 * 1024 or len(archive.infolist()) > 2000:
                    raise DocumentError("Word 文件解压后过大")
            document = Document(BytesIO(content))
            sections = [paragraph.text for paragraph in document.paragraphs]
            sections.extend(
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables for row in table.rows
            )
            text = "\n".join(sections)
        except (BadZipFile, PackageNotFoundError, KeyError, XMLSyntaxError) as error:
            raise DocumentError("无法读取 DOCX，请确认文件有效且不是旧版 .doc") from error
    elif extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise DocumentError("文件内容不是有效的 PDF")
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise DocumentError("暂不支持加密 PDF，请先移除文件密码")
            if len(reader.pages) > 30:
                raise DocumentError("简历 PDF 最多支持 30 页")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except PdfReadError as error:
            raise DocumentError("无法读取 PDF，请重新导出") from error
    else:
        raise DocumentError("仅支持 .docx 和 .pdf；旧版 .doc 请另存为 .docx")
    text = text.strip()
    if not text:
        raise DocumentError("未读取到文字；扫描版 PDF 请先做 OCR 或上传 Word")
    if len(text) > MAX_TEXT_LENGTH:
        raise DocumentError("简历文字超过 80,000 字符，请精简后重新上传")
    return text


def resume_sections(profile: ResumeProfile) -> tuple[tuple[str, Sequence[Model]], ...]:
    return (
        ("Experience", profile.experience), ("Education", profile.education),
        ("Publications", profile.publications),
        ("Projects", profile.projects), ("Certifications", profile.certifications),
        ("Languages", profile.languages),
    )


def publication_details(
    entry: Publication, *, include_date: bool = True, include_authors: bool = True,
) -> str:
    return " | ".join(value for value in (
        ", ".join(entry.authors) if include_authors else "", entry.venue, entry.date if include_date else "",
        "DOI: " + entry.doi if entry.doi else "", entry.url,
    ) if value)


def resume_entry_text(entry: Model) -> tuple[str, str, list[str]]:
    if isinstance(entry, Publication):
        return entry.title, publication_details(entry), entry.details
    data = entry.model_dump()
    heading = " | ".join(
        data[key] for key in (
            "title", "company", "institution", "degree", "field",
            "name", "issuer", "proficiency",
        ) if data.get(key)
    )
    detail = " | ".join(
        data[key] for key in ("location", "start_date", "end_date", "date", "url")
        if data.get(key)
    )
    return heading, detail, data.get("bullets", data.get("details", []))


def resume_docx(profile: ResumeProfile) -> BytesIO:
    document = Document()
    contact = profile.contact
    document.add_heading(contact.name or "Resume", 0)
    document.add_paragraph(" | ".join(
        value for key, value in contact.model_dump().items() if key != "name" and value
    ))
    if profile.summary:
        document.add_heading("Summary", 1)
        document.add_paragraph(profile.summary)
    if profile.skills:
        document.add_heading("Skills", 1)
        document.add_paragraph(", ".join(profile.skills))
    for title, entries in resume_sections(profile):
        if not entries:
            continue
        document.add_heading(title, 1)
        for entry in entries:
            heading, detail, bullets = resume_entry_text(entry)
            if heading:
                document.add_heading(heading, 2)
            if detail:
                document.add_paragraph(detail)
            for bullet in bullets:
                document.add_paragraph(bullet, style="List Bullet")
    result = BytesIO()
    document.save(result)
    result.seek(0)
    return result
