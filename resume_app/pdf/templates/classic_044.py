"""044: centered identity, gray section bars and a date/content column layout."""

from collections.abc import Sequence
from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from docx.styles.style import ParagraphStyle as WordParagraphStyle
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Flowable, Table, TableStyle

from resume_app.documents import DocumentError, publication_details, resume_entry_text
from resume_app.models import Model, Publication, ResumeProfile
from resume_app.pdf.rendering import PdfFonts, paragraph, validate_profile, validate_text
from resume_app.pdf.templates.base import ResumeTemplate


class Classic044Template(ResumeTemplate):
    id = "classic-044"
    name = "经典灰条 · 044"
    description = "居中姓名、灰色章节条、日期与经历分栏；自动续页，无需 Word 转换。"
    version = "reportlab-044-v1"
    reference_document = "references/044.docx"
    page_size = A4
    left_margin = right_margin = 31.2
    top_margin = 28.35
    bottom_margin = 34.0
    serif = True
    uses_job_title = True
    name_size = 28
    section_size = 14
    body_size = 11
    body_leading = 13.2
    date_width = 105.0
    section_fill = "D9D9D9"

    def sections(self, profile: ResumeProfile) -> tuple[tuple[str, str | Sequence[Model]], ...]:
        return (
            ("PROFILE", profile.summary), ("EXPERIENCE", profile.experience),
            ("EDUCATION", profile.education), ("SKILLS", ", ".join(profile.skills)),
            ("PUBLICATIONS", profile.publications),
            ("PROJECTS", profile.projects), ("CERTIFICATIONS", profile.certifications),
            ("LANGUAGES", profile.languages),
        )

    def entry_text(self, entry: Model) -> tuple[str, str, str, list[str]]:
        if isinstance(entry, Publication):
            return entry.title, publication_details(entry, include_date=False), entry.date, entry.details
        data = entry.model_dump()
        heading, _, bullets = resume_entry_text(entry)
        dates = " - ".join(data[key] for key in ("start_date", "end_date", "date") if data.get(key))
        detail = " | ".join(data[key] for key in ("location", "url") if data.get(key))
        return heading, detail, dates, bullets

    def story(self, profile: ResumeProfile, fonts: PdfFonts, job_title: str) -> list[Flowable]:
        body = ParagraphStyle(
            "Body044", fontName=fonts.regular.fontName, fontSize=self.body_size,
            leading=self.body_leading, textColor=HexColor("#404040"),
            spaceAfter=4, wordWrap="CJK",
        )
        name_style = ParagraphStyle(
            "Name044", parent=body, fontSize=self.name_size, leading=32,
            textColor=HexColor("#000000"), alignment=TA_CENTER,
            spaceAfter=2, keepWithNext=True,
        )
        target_style = ParagraphStyle(
            "Target044", parent=body, fontSize=10, leading=12,
            alignment=TA_CENTER, spaceAfter=8, keepWithNext=True,
        )
        contact_style = ParagraphStyle(
            "Contact044", parent=body, alignment=TA_CENTER,
            spaceAfter=10, keepWithNext=True,
        )
        section_style = ParagraphStyle(
            "Section044", parent=body, fontSize=self.section_size, leading=18,
            textColor=HexColor("#000000"), backColor=HexColor("#" + self.section_fill),
            borderPadding=3, spaceBefore=10, spaceAfter=5, keepWithNext=True,
        )
        entry_style = ParagraphStyle("Entry044", parent=body, keepWithNext=True)
        bullet_style = ParagraphStyle(
            "Bullet044", parent=body, leftIndent=15, bulletIndent=3,
            bulletFontName=fonts.regular.fontName, bulletFontSize=self.body_size,
        )
        story: list[Flowable] = [
            paragraph(profile.contact.name, name_style, fonts, bold=True),
        ]
        if job_title:
            story.append(paragraph("TARGET: " + job_title, target_style, fonts))
        contact = " | ".join(
            value for key, value in profile.contact.model_dump().items()
            if key != "name" and value
        )
        if contact:
            story.append(paragraph(contact, contact_style, fonts))
        for title, content in self.sections(profile):
            if not content:
                continue
            story.append(paragraph(title, section_style, fonts, bold=True))
            if isinstance(content, str):
                story.append(paragraph(content, body, fonts))
                continue
            for entry in content:
                heading, detail, dates, bullets = self.entry_text(entry)
                entry_story: list[Flowable] = []
                if heading:
                    entry_story.append(paragraph(heading, entry_style, fonts, bold=True))
                if detail:
                    entry_story.append(paragraph(detail, body, fonts))
                entry_story.extend(
                    paragraph(bullet, bullet_style, fonts, bullet="\u2022")
                    for bullet in bullets if bullet
                )
                story.append(Table(
                    [[paragraph(dates, body, fonts, bold=True), entry_story]],
                    colWidths=[self.date_width, self.content_width - self.date_width],
                    hAlign="LEFT", splitInRow=1,
                    style=TableStyle([
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]),
                ))
        return story

    def render_docx(self, profile: ResumeProfile, *, job_title: str = "") -> bytes:
        validate_profile(profile)
        validate_text(job_title)
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Twips(11906), Twips(16838)
        section.left_margin = section.right_margin = Pt(self.left_margin)
        section.top_margin, section.bottom_margin = Pt(self.top_margin), Pt(self.bottom_margin)

        def style(name: str, size: int, *, bold: bool = False) -> WordParagraphStyle:
            result = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            if not isinstance(result, WordParagraphStyle):
                raise DocumentError("无法创建 Word 段落样式")
            result.font.name, result.font.size, result.font.bold = "Garamond", Pt(size), bold
            result.paragraph_format.line_spacing = 1.2
            result.paragraph_format.space_after = Pt(4)
            fonts = result.element.find(".//" + qn("w:rFonts"))
            if fonts is not None:
                fonts.set(qn("w:eastAsia"), "SimSun")
            return result

        body = style("Resume Body", self.body_size)
        name_style = style("Resume Name", self.name_size, bold=True)
        name_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_style.paragraph_format.keep_with_next = True
        target = style("Resume Target", 10)
        target.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        target.paragraph_format.keep_with_next = True
        contact_style = style("Resume Contact", self.body_size)
        contact_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_style.paragraph_format.keep_with_next = True
        contact_style.paragraph_format.space_after = Pt(10)
        heading_style = style("Resume Section", self.section_size, bold=True)
        heading_style.paragraph_format.keep_with_next = True
        heading_style.paragraph_format.space_before = Pt(10)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), self.section_fill)
        heading_style.element.get_or_add_pPr().append(shade)
        entry_style = style("Resume Entry", self.body_size, bold=True)
        entry_style.paragraph_format.keep_with_next = True
        date_style = style("Resume Date", self.body_size, bold=True)
        bullet_style = style("Resume Bullet", self.body_size)
        bullet_style.base_style = document.styles["List Bullet"]
        document.add_paragraph(profile.contact.name, name_style)
        if job_title:
            document.add_paragraph("TARGET: " + job_title, target)
        contact = " | ".join(
            value for key, value in profile.contact.model_dump().items()
            if key != "name" and value
        )
        if contact:
            document.add_paragraph(contact, contact_style)
        for title, content in self.sections(profile):
            if not content:
                continue
            document.add_paragraph(title, heading_style)
            if isinstance(content, str):
                document.add_paragraph(content, body)
                continue
            for entry in content:
                heading, detail, dates, bullets = self.entry_text(entry)
                table = document.add_table(rows=1, cols=2)
                table.autofit = False
                widths = (Pt(self.date_width), Twips(10658 - round(self.date_width * 20)))
                for column, cell, width in zip(table.columns, table.rows[0].cells, widths):
                    column.width = width
                    cell.width = width
                left, right = table.rows[0].cells
                left.paragraphs[0].text, left.paragraphs[0].style = dates, date_style
                right.paragraphs[0].text, right.paragraphs[0].style = heading, entry_style
                if detail:
                    right.add_paragraph(detail, body)
                for bullet in bullets:
                    if bullet:
                        right.add_paragraph(bullet, bullet_style)
        document.core_properties.title = ("Resume - " + profile.contact.name)[:255]
        document.core_properties.author = profile.contact.name[:255]
        result = BytesIO()
        document.save(result)
        return result.getvalue()
