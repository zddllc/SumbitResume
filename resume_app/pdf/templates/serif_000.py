"""000: a compact serif resume with an identity rule and underlined headings."""

from collections.abc import Sequence
from io import BytesIO
import os
from pathlib import Path
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from docx.styles.style import ParagraphStyle as WordParagraphStyle
from reportlab.lib.colors import black, white
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Flowable, Paragraph, Table, TableStyle

from resume_app.documents import DocumentError, publication_details, resume_entry_text
from resume_app.models import Education, Experience, Model, Publication, ResumeProfile
from resume_app.pdf.rendering import PdfFonts, paragraph, registered_font, validate_profile, validate_text
from resume_app.pdf.templates.base import ResumeTemplate


class ArrowBullet(Paragraph):
    draw_arrow = True

    def draw(self) -> None:
        super().draw()
        if not self.draw_arrow:
            return
        canvas = self.canv
        canvas.saveState()
        canvas.setFillColor(black)
        canvas.translate(0, self.height - self.style.fontSize + 0.5)
        arrow = canvas.beginPath()
        arrow.moveTo(0, 0)
        arrow.lineTo(8, 3.5)
        arrow.lineTo(0, 7)
        arrow.lineTo(2.5, 3.5)
        arrow.close()
        canvas.drawPath(arrow, stroke=0, fill=1)
        canvas.setFillColor(white)
        highlight = canvas.beginPath()
        highlight.moveTo(0.8, 6.2)
        highlight.lineTo(6.8, 3.7)
        highlight.lineTo(2.8, 3.7)
        highlight.close()
        canvas.drawPath(highlight, stroke=0, fill=1)
        canvas.restoreState()

    def split(self, availWidth: float, availHeight: float) -> list[Flowable]:
        parts = super().split(availWidth, availHeight)
        for index, part in enumerate(parts):
            if isinstance(part, ArrowBullet):
                part.draw_arrow = self.draw_arrow and index == 0
        return parts


def role_fonts(fallback: PdfFonts) -> PdfFonts:
    windows = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    candidates = (
        [(windows / regular, windows / italic) for regular, italic in (
            ("calibri.ttf", "calibrii.ttf"), ("arial.ttf", "ariali.ttf"),
        )] if sys.platform == "win32" else [
            (Path(root) / "LiberationSans-Regular.ttf", Path(root) / "LiberationSans-Italic.ttf")
            for root in ("/usr/share/fonts/truetype/liberation2", "/usr/share/fonts/truetype/liberation")
        ] + [(Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
              Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"))]
    )
    for regular, italic in candidates:
        if regular.is_file() and italic.is_file():
            return PdfFonts(
                registered_font(str(regular)), registered_font(str(regular)),
                fallback.fallback, registered_font(str(italic)),
            )
    return PdfFonts(fallback.fallback, fallback.fallback, fallback.fallback)


class Serif000Template(ResumeTemplate):
    id = "serif-000"
    name = "衬线履历 · 000"
    description = "Garamond 正文、下划线章节标题、顶部细横线及紧凑论文排版。"
    version = "reportlab-000-v2"
    reference_document = "references/000.docx"
    page_size = A4
    # Compensate for the shared ReportLab frame's six-point inner padding.
    left_margin = right_margin = 54 - 6
    top_margin = bottom_margin = 72 - 6
    serif = True
    name_size = 14
    body_size = 11.5
    section_size = 15
    publication_size = 10

    def experience_meta(self, entry: Experience) -> tuple[str, str]:
        return (
            " | ".join(value for value in (entry.company, entry.location) if value),
            " - ".join(value for value in (entry.start_date, entry.end_date) if value),
        )

    def sections(self, profile: ResumeProfile) -> tuple[tuple[str, str | Sequence[Model]], ...]:
        return (
            ("Industry Experience", profile.experience), ("Education", profile.education),
            ("Publications", profile.publications), ("Technical Skills", ", ".join(profile.skills)),
            ("Projects", profile.projects), ("Certifications", profile.certifications),
            ("Languages", profile.languages),
        )

    def entry_text(self, entry: Model) -> tuple[str, str, list[str]]:
        if isinstance(entry, Experience):
            company, dates = self.experience_meta(entry)
            return entry.title, " | ".join(value for value in (company, dates) if value), entry.bullets
        if isinstance(entry, Education):
            return (
                " | ".join(value for value in (entry.degree, entry.field) if value),
                " | ".join(value for value in (entry.institution, entry.start_date, entry.end_date) if value),
                entry.details,
            )
        if isinstance(entry, Publication):
            citation = ". ".join(value for value in (", ".join(entry.authors), entry.title) if value)
            detail = publication_details(entry, include_authors=False)
            return "", " | ".join(value for value in (citation, detail) if value), entry.details
        return resume_entry_text(entry)

    def story(self, profile: ResumeProfile, fonts: PdfFonts, job_title: str) -> list[Flowable]:
        roles = role_fonts(fonts)
        body = ParagraphStyle(
            "Body000", fontName=fonts.regular.fontName, fontSize=self.body_size,
            leading=14, textColor=black, spaceAfter=4, wordWrap="CJK",
        )
        name_style = ParagraphStyle(
            "Name000", parent=body, fontSize=self.name_size, leading=17,
            spaceAfter=3, keepWithNext=True,
        )
        contact_style = ParagraphStyle("Contact000", parent=body, fontSize=12, leading=14)
        right_contact = ParagraphStyle("RightContact000", parent=contact_style, alignment=TA_RIGHT)
        section_style = ParagraphStyle(
            "Section000", parent=body, fontName="Times-Bold",
            fontSize=self.section_size, leading=19,
            spaceBefore=9, spaceAfter=5, keepWithNext=True,
        )
        role_style = ParagraphStyle(
            "Role000", parent=body, fontName=roles.regular.fontName, fontSize=12, leading=15,
            spaceBefore=5, spaceAfter=2, keepWithNext=True,
        )
        detail_style = ParagraphStyle("Detail000", parent=body, keepWithNext=True)
        company_style = ParagraphStyle("Company000", parent=body, spaceAfter=0)
        date_style = ParagraphStyle("Date000", parent=company_style, alignment=TA_RIGHT)
        bullet_style = ParagraphStyle(
            "Bullet000", parent=body, leftIndent=21, bulletIndent=0,
        )
        citation_style = ParagraphStyle(
            "Citation000", parent=body, fontSize=self.publication_size, leading=12,
            spaceBefore=3, keepWithNext=False,
        )
        contact = profile.contact
        rows = [[paragraph(contact.name, name_style, fonts, bold=True),
                 paragraph(contact.email, right_contact, fonts)]]
        links = "\n".join(value for value in (contact.linkedin, contact.website) if value)
        location = " | ".join(value for value in (contact.phone, contact.location) if value)
        if links or location:
            rows.append([paragraph(links, contact_style, fonts), paragraph(location, right_contact, fonts)])
        header = Table(
            rows, colWidths=[self.content_width * 0.55, self.content_width * 0.45],
            hAlign="LEFT", splitInRow=1,
            style=TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LINEBELOW", (0, -1), (-1, -1), 0.5, black),
            ]),
        )
        story: list[Flowable] = [header]
        if profile.summary:
            story.append(paragraph(profile.summary, body, fonts))
        for title, content in self.sections(profile):
            if not content:
                continue
            story.append(Paragraph(f"<u>{title}</u>", section_style))
            if isinstance(content, str):
                story.append(paragraph(content, body, fonts))
                continue
            for entry in content:
                heading, detail, bullets = self.entry_text(entry)
                is_publication = isinstance(entry, Publication)
                if heading:
                    story.append(paragraph(
                        heading, role_style, roles if isinstance(entry, Experience) else fonts,
                        italic=isinstance(entry, Experience), bold=not isinstance(entry, Experience),
                    ))
                if isinstance(entry, Experience):
                    company, dates = self.experience_meta(entry)
                    if company or dates:
                        company_text = " | ".join(value for value in (
                            fonts.markup(entry.company, bold=True), fonts.markup(entry.location),
                        ) if value)
                        metadata = Table(
                            [[Paragraph(company_text, company_style), paragraph(dates, date_style, fonts)]],
                            colWidths=[self.content_width * 0.65, self.content_width * 0.35],
                            hAlign="LEFT", splitInRow=1,
                            style=TableStyle([
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                ("RIGHTPADDING", (0, 0), (0, -1), 12),
                                ("RIGHTPADDING", (1, 0), (1, -1), 0),
                                ("TOPPADDING", (0, 0), (-1, -1), 0),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                            ]),
                        )
                        metadata.keepWithNext = bool(bullets)
                        story.append(metadata)
                elif detail:
                    story.append(paragraph(detail, citation_style if is_publication else detail_style, fonts))
                for bullet in bullets:
                    if bullet:
                        story.append(
                            paragraph(bullet, citation_style, fonts) if is_publication
                            else ArrowBullet(fonts.markup(bullet), bullet_style)
                        )
        return story

    def render_docx(self, profile: ResumeProfile, *, job_title: str = "") -> bytes:
        validate_profile(profile)
        validate_text(job_title)
        document = Document()
        page = document.sections[0]
        page.page_width, page.page_height = Twips(11906), Twips(16838)
        page.left_margin = page.right_margin = Pt(54)
        page.top_margin = page.bottom_margin = Pt(72)

        def style(name: str, size: float, *, bold: bool = False, italic: bool = False) -> WordParagraphStyle:
            result = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            if not isinstance(result, WordParagraphStyle):
                raise DocumentError("无法创建 Word 段落样式")
            result.font.name, result.font.size = "Garamond", Pt(size)
            result.font.bold, result.font.italic = bold, italic
            result.paragraph_format.space_after = Pt(4)
            result.paragraph_format.line_spacing = 1.05
            result.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fonts = result.element.find(".//" + qn("w:rFonts"))
            if fonts is not None:
                fonts.set(qn("w:eastAsia"), "SimSun")
            return result

        body = style("Resume000 Body", self.body_size)
        name_style = style("Resume000 Name", self.name_size, bold=True)
        name_style.font.small_caps = True
        contact_style = style("Resume000 Contact", 12)
        section_style = style("Resume000 Section", self.section_size, bold=True)
        section_style.font.name = "Times New Roman"
        section_style.font.underline = WD_UNDERLINE.SINGLE
        section_style.paragraph_format.line_spacing = Pt(19)
        section_style.paragraph_format.space_before = Pt(9)
        section_style.paragraph_format.keep_with_next = True
        role_style = style("Resume000 Role", 12, italic=True)
        role_style.font.name = "Calibri"
        role_style.paragraph_format.keep_with_next = True
        entry_style = style("Resume000 Entry", self.body_size, bold=True)
        entry_style.paragraph_format.keep_with_next = True
        detail_style = style("Resume000 Detail", self.body_size)
        detail_style.paragraph_format.keep_with_next = True
        company_style = style("Resume000 Company", self.body_size)
        company_style.paragraph_format.keep_with_next = True
        date_style = style("Resume000 Date", self.body_size)
        date_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        date_style.paragraph_format.keep_with_next = True
        bullet_style = style("Resume000 Bullet", self.body_size)
        bullet_style.paragraph_format.left_indent = Pt(21)
        bullet_style.paragraph_format.first_line_indent = Pt(-21)
        citation_style = style("Resume000 Citation", self.publication_size)

        numbering = document.part.numbering_part.element
        abstract_id = max(int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in (("start", "1"), ("numFmt", "bullet"), ("lvlText", "\uf0d8"), ("lvlJc", "left")):
            node = OxmlElement("w:" + tag)
            node.set(qn("w:val"), value)
            level.append(node)
        properties = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "420")
        indent.set(qn("w:hanging"), "420")
        properties.append(indent)
        level.append(properties)
        run_properties = OxmlElement("w:rPr")
        marker_fonts = OxmlElement("w:rFonts")
        for attribute in ("ascii", "hAnsi"):
            marker_fonts.set(qn("w:" + attribute), "Wingdings")
        run_properties.append(marker_fonts)
        level.append(run_properties)
        abstract.append(level)
        numbering.insert_element_before(abstract, "w:num")
        number_id = numbering.add_num(abstract_id).numId

        contact = profile.contact
        table = document.add_table(rows=2, cols=2)
        table.autofit = False
        for index, ratio in enumerate((0.55, 0.45)):
            width = Pt((page.page_width.pt - 108) * ratio)
            table.columns[index].width = width
            for cell in table.columns[index].cells:
                cell.width = width
        table.cell(0, 0).paragraphs[0].text = contact.name
        table.cell(0, 0).paragraphs[0].style = name_style
        table.cell(0, 1).paragraphs[0].text = contact.email
        table.cell(1, 0).paragraphs[0].text = "\n".join(value for value in (contact.linkedin, contact.website) if value)
        table.cell(1, 1).paragraphs[0].text = " | ".join(value for value in (contact.phone, contact.location) if value)
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                for item in cell.paragraphs:
                    if item.style != name_style:
                        item.style = contact_style
                if index == 1:
                    for item in cell.paragraphs:
                        item.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        borders = OxmlElement("w:tblBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:color"), "000000")
        borders.append(bottom)
        table._tbl.tblPr.append(borders)

        if profile.summary:
            document.add_paragraph(profile.summary, body)
        for title, content in self.sections(profile):
            if not content:
                continue
            document.add_paragraph(title, section_style)
            if isinstance(content, str):
                document.add_paragraph(content, body)
                continue
            for entry in content:
                heading, detail, bullets = self.entry_text(entry)
                is_publication = isinstance(entry, Publication)
                if heading:
                    document.add_paragraph(heading, role_style if isinstance(entry, Experience) else entry_style)
                if isinstance(entry, Experience):
                    company, dates = self.experience_meta(entry)
                    if company or dates:
                        metadata = document.add_table(rows=1, cols=2)
                        metadata.autofit = False
                        for index, ratio in enumerate((0.65, 0.35)):
                            width = Pt((page.page_width.pt - 108) * ratio)
                            metadata.columns[index].width = width
                            metadata.cell(0, index).width = width
                        left, right = metadata.rows[0].cells
                        left.paragraphs[0].style = company_style
                        left.paragraphs[0].add_run(entry.company).bold = True
                        if entry.location:
                            left.paragraphs[0].add_run((" | " if entry.company else "") + entry.location)
                        right.paragraphs[0].text = dates
                        right.paragraphs[0].style = date_style
                        margins = OxmlElement("w:tblCellMar")
                        for side, width in (("top", 0), ("bottom", 0), ("left", 0), ("right", 0)):
                            margin = OxmlElement("w:" + side)
                            margin.set(qn("w:w"), str(width))
                            margin.set(qn("w:type"), "dxa")
                            margins.append(margin)
                        metadata._tbl.tblPr.append(margins)
                elif detail:
                    document.add_paragraph(detail, citation_style if is_publication else detail_style)
                for bullet in bullets:
                    if bullet:
                        item = document.add_paragraph(bullet, citation_style if is_publication else bullet_style)
                        if not is_publication:
                            num = item._p.get_or_add_pPr().get_or_add_numPr()
                            num.get_or_add_ilvl().val = 0
                            num.get_or_add_numId().val = number_id
        document.core_properties.title = ("Resume - " + contact.name)[:255]
        document.core_properties.author = contact.name[:255]
        document.core_properties.last_modified_by = ""
        result = BytesIO()
        document.save(result)
        return result.getvalue()
